from pathlib import Path
import io, json, re
import numpy as np, pandas as pd
import plotly.graph_objects as go
import streamlit as st
from scripts.universal_data import universe,get_company,get_snapshot,entity_history,peer_snapshot,peer_metric_history,industry_snapshot,industry_metric_history,industry_label,period_date,num,coverage
from scripts.multisector_valuation import valuation, latest_market_price
from scripts.securities_public_metrics import latest_capital_ratio
from scripts.multisector_report import generate_docx,generate_pdf
from scripts.sector_templates import get_template
from scripts.sector_kpi_engine import sector_kpi_table
from scripts.intelligent_analyst import analyze as intelligent_analyze
from scripts.valuation_regime import assess as valuation_regime
from scripts.three_methodology_rating import rate_company as rate_three_methodologies
from scripts.fair_value_range import fair_value_range
from scripts.rating_committee_engine import committee_pack
from scripts.valuation_triangulation import triangulate
from scripts.rating_evidence_engine import rating_evidence
try:
    from scripts.coverage_engine import build_coverage_matrix
except Exception:
    build_coverage_matrix=None
try: from scripts.credit_rating_engine import build_credit_rating
except Exception: build_credit_rating=None

ROOT=Path(__file__).resolve().parent; DATA=ROOT/'data'
BUILD_TAG='V8.121.0-INSURANCE-PEERS'

from scripts.company_profile import load_company_profile, profile_overview_text, profile_source_text

def _override_peer_info(ticker):
    """Read analyst peer override directly at UI boundary so Cloud/local resolve identically."""
    path=ROOT/'config'/'dynamic_peer_overrides.csv'
    try:
        o=pd.read_csv(path)
        if o.empty or not {'TargetTicker','PeerTicker'}.issubset(o.columns): return None, []
        o=o.copy(); o['TargetTicker']=o['TargetTicker'].astype(str).str.upper().str.strip(); o['PeerTicker']=o['PeerTicker'].astype(str).str.upper().str.strip()
        if 'Active' in o.columns: o=o[pd.to_numeric(o['Active'],errors='coerce').fillna(1).eq(1)]
        z=o[o['TargetTicker'].eq(str(ticker).upper().strip())].copy()
        if z.empty:return None, []
        if 'Priority' in z.columns:
            z['_p']=pd.to_numeric(z['Priority'],errors='coerce').fillna(9999); z=z.sort_values(['_p','PeerTicker'])
        peers=[x for x in z.PeerTicker.astype(str).tolist() if x and x!=str(ticker).upper().strip()]
        return 'Nhóm peer chuyên ngành (analyst override)', peers
    except Exception:
        return None, []

def _final_chart_guard(metric, df):
    """Last-mile validity guard: malformed upstream values must never reach Plotly."""
    if df is None or not len(df): return df
    q=df.copy(); m=str(metric)
    for c in ['IndustryMean','IndustryMedian']:
        if c not in q.columns: continue
        v=pd.to_numeric(q[c],errors='coerce')
        if m=='DebtEquity': v=v.where((v>=0)&(v<=10))
        elif m=='CurrentRatio': v=v.where((v>0)&(v<=20))
        elif m=='PE': v=v.where((v>0)&(v<=200))
        elif m=='PB': v=v.where((v>0)&(v<=20))
        elif m in {'ROE','ROA'}: v=v.where((v>=-2)&(v<=2))
        q[c]=v
    keep=[c for c in ['IndustryMean','IndustryMedian'] if c in q.columns]
    return q.dropna(subset=keep,how='all') if keep else q


def _security_raw_latest_value(ticker, metric):
    """Latest actual Securities value from the same raw semantic files used by refresh.
    This is a UI fallback only; it never fabricates observations.
    """
    h=_security_raw_metric_history(ticker,metric)
    if h is not None and len(h):
        v=pd.to_numeric(h['Value'],errors='coerce').dropna()
        if len(v): return float(v.iloc[-1])
    return None

def _snapshot_row(ticker):
    return get_snapshot(ticker)

def _securities_peer_frame(ticker, max_peers=10):
    """Build a dynamic Securities peer set when industry_snapshot is empty.
    Priority: analyst override -> closest listed securities firms by size.
    """
    t=str(ticker).upper().strip()
    try:
        q=industry_snapshot(t)
        if q is not None and len(q):
            q=q.copy()
            if 'Ticker' in q.columns:
                q['Ticker']=q['Ticker'].astype(str).str.upper().str.strip()
                if q.Ticker.nunique()>=2:
                    # Preserve upstream peer selection, but rebuild each row through the
                    # enriched snapshot/raw fallback so sparse industry_snapshot columns
                    # cannot blank all charts and KPI tables.
                    _ticks=[t]+[x for x in q['Ticker'].tolist() if x!=t][:max_peers]
                    return pd.DataFrame([_snapshot_row(x) for x in _ticks]).drop_duplicates('Ticker',keep='first')
    except Exception:
        pass

    # Analyst override, if configured.
    try:
        _, override_peers=_override_peer_info(t)
    except Exception:
        override_peers=[]
    peer_tickers=[x for x in override_peers if str(x).upper().strip()!=t]

    # Otherwise build dynamically from all listed/registered securities companies.
    if not peer_tickers:
        try:
            uu=universe().copy()
            uu=uu[uu.EntityType.astype(str).str.upper().eq('SECURITIES')]
            cand=[str(x).upper().strip() for x in uu.Ticker.tolist() if str(x).upper().strip()!=t]
        except Exception:
            cand=[]
        target=_snapshot_row(t)
        target_size=None
        for k in ('TotalAssets','Equity','Revenue','MarketCap'):
            try:
                v=float(target.get(k))
                if np.isfinite(v) and v>0:
                    target_size=(k,v); break
            except Exception:
                pass
        scored=[]
        for x in cand:
            r=_snapshot_row(x)
            score=999.0
            if target_size:
                k,tv=target_size
                try:
                    pv=float(r.get(k))
                    if np.isfinite(pv) and pv>0:
                        score=abs(np.log(pv/tv))
                except Exception:
                    pass
            scored.append((score,x,r))
        scored.sort(key=lambda z:(z[0],z[1]))
        peer_tickers=[x for _,x,_ in scored[:max_peers]]

    rows=[_snapshot_row(t)]+[_snapshot_row(x) for x in peer_tickers[:max_peers]]
    return pd.DataFrame(rows).drop_duplicates('Ticker',keep='first')

def _peer_frame_safe(ticker):
    try:
        et=str(get_company(ticker).get('EntityType','')).upper()
    except Exception:
        et=''
    if et=='SECURITIES':
        return _securities_peer_frame(ticker,10)
    try:
        return industry_snapshot(ticker)
    except Exception:
        return pd.DataFrame()

def _securities_kpi_fallback(ticker, peer_df):
    """Create the same benchmark table from actual target/peer snapshots when
    sector_kpi_table has no usable securities observations.
    """
    if peer_df is None or not len(peer_df) or 'Ticker' not in peer_df.columns:
        return pd.DataFrame()
    t=str(ticker).upper().strip()
    q=peer_df.copy()
    q['Ticker']=q['Ticker'].astype(str).str.upper().str.strip()
    target=q[q.Ticker.eq(t)]
    peers=q[~q.Ticker.eq(t)]
    specs=[
        ('Vốn & đòn bẩy','AvailableCapitalRatio','Tỷ lệ an toàn vốn khả dụng'),
        ('Vốn & đòn bẩy','DebtEquity','Nợ/VCSH'),
        ('Khả năng sinh lợi & hiệu quả','ROE','ROE/ROAE'),
        ('Khả năng sinh lợi & hiệu quả','ROA','ROA/ROAA'),
        ('Nguồn vốn & thanh khoản','CurrentRatio','Hệ số thanh toán hiện hành'),
        ('Định giá','PB','P/B'),
        ('Định giá','PE','P/E'),
    ]
    rows=[]
    for grp,m,label in specs:
        if m not in q.columns: vals=pd.Series(dtype=float)
        else: vals=pd.to_numeric(peers[m],errors='coerce').dropna()
        tv=None
        if len(target) and m in target.columns:
            vv=pd.to_numeric(target[m],errors='coerce').dropna()
            if len(vv): tv=float(vv.iloc[-1])
        mean=float(vals.mean()) if len(vals) else None
        med=float(vals.median()) if len(vals) else None
        rows.append({
            'Metric':m,'Nhóm phân tích':grp,'Chỉ tiêu':label,
            'Doanh nghiệp':tv,'Trung bình ngành':mean,'Trung vị ngành':med,
            'Số DN có dữ liệu':int(len(vals)),
            'Chênh lệch với TB ngành':(tv-mean) if tv is not None and mean is not None else None,
            'Trạng thái dữ liệu':('Có dữ liệu' if tv is not None else 'N/A – cần bổ sung nguồn doanh nghiệp')
        })
    return pd.DataFrame(rows)

def _dynamic_peer_label_ui(ticker, fallback):
    # V8.73.1: analyst override wins at the UI boundary, independent of imported-module cache.
    src, peers=_override_peer_info(ticker)
    if src and peers:
        return f'Nhóm tương đồng động: {len(peers)} DN từ {src}'
    try:
        from scripts.dynamic_peer_engine import dynamic_peer_label
        x=dynamic_peer_label(ticker)
        return x if x else fallback
    except Exception:
        return fallback

def _rating_no_vn(value):
    """Display rating grades without the historical 'vn' prefix."""
    if value is None or str(value) in ('N/A','nan','None'):
        return 'Chưa thẩm định'
    s=str(value).strip()
    if re.match(r'(?i)^vn(?=[A-C])',s):
        s=s[2:]
    return s.upper() if re.match(r'(?i)^[A-C]{1,3}(?:[+-])?(?:-[A-C]{1,3})?$',s) else s

# V8.94.1: unique Plotly key per render, including repeated calls from loops.
_plotly_seq = 0
def _plotly_chart_safe(fig, **kwargs):
    global _plotly_seq
    if fig is None:
        return None
    try:
        if hasattr(fig,'data') and len(fig.data)==0:
            return None
    except Exception:
        pass
    _plotly_seq += 1
    # Streamlit 1.5x deprecates use_container_width in favor of width.
    if "use_container_width" in kwargs:
        ucw = kwargs.pop("use_container_width")
        kwargs.setdefault("width", "stretch" if ucw else "content")
    ticker_key = str(globals().get("selected", globals().get("t", ""))).upper().strip() or "NA"
    kwargs["key"] = f"plotly_{ticker_key}_{_plotly_seq}"
    return st.plotly_chart(fig, **kwargs)

st.set_page_config(page_title='Nền tảng Phân tích, Định giá, M&A & XHTN Doanh nghiệp Việt Nam',page_icon='🏢',layout='wide')
st.markdown('''<style>
.block-container{padding-top:4.35rem!important;max-width:1550px}.muted{color:#9CA3AF;font-size:.86rem}
[data-testid="stSidebar"]{min-width:300px;max-width:300px}
h1{font-size:1.95rem!important;line-height:1.35!important;overflow:visible!important}
h2{font-size:1.4rem!important} h3{font-size:1.1rem!important}
.platform-title{
    font-size:1.82rem;font-weight:800;line-height:1.48!important;
    padding:.52rem 0 .20rem 0!important;margin:0 0 .35rem 0!important;
    overflow:visible!important;white-space:normal;word-break:normal;
}
.platform-disclaimer{
    font-size:.86rem;line-height:1.55!important;color:#AEB4BE;
    padding:.08rem 0 .18rem 0!important;margin:0 0 1.00rem 0!important;
    overflow:visible!important;white-space:normal;
}
[data-testid="stMetricValue"]{font-size:1.3rem} div[data-testid="stTabs"] button{white-space:nowrap;font-weight:650}
</style>''',unsafe_allow_html=True)

def vi(x,d=1):
    v=num(x)
    if v is None:return 'N/A'
    return f'{v:,.{d}f}'.replace(',','X').replace('.',',').replace('X','.')
def pct(x):return 'N/A' if num(x) is None else vi(num(x)*100,1)+'%'
def mult(x):return 'N/A' if num(x) is None else vi(x,2)+'x'
def _price_vnd(x):
    v=num(x)
    if v is None:return None
    v=float(v)
    if 0<v<500:v*=1000.0
    while v>1_000_000:v/=1000.0
    return v if 100<=v<=1_000_000 else None
def money(x):
    v=_price_vnd(x)
    return 'Thiếu dữ liệu' if v is None else vi(v,0)+' đồng/cp'
def bn(x):return 'N/A' if num(x) is None else vi(num(x)/1e9,0)+' tỷ đồng'


def _security_raw_metric_history(ticker, metric):
    """Fallback to Vnstock raw semantic files for Securities trend metrics.
    Returns canonical Period/Value history without API calls.
    """
    t=str(ticker).upper().strip()
    metric_ids={
        'ROE':('ratio',['RT_PRT_ROE']),
        'ROA':('ratio',['RT_PRT_ROA']),
        'DebtEquity':('ratio',['RT_LEV_DE']),
        'CurrentRatio':('ratio',['RT_LQD_CR']),
        'PB':('ratio',['RT_VALUE_PB']),
        'PE':('ratio',['RT_VALUE_PE']),
        'Revenue':('income',['IS_NET_REVENUE','IS_REVENUE','IS_OPERATING_REVENUE']),
        'NPAT':('income',['IS_NET_PROFIT_AFTER_TAX','IS_NET_PROFIT','IS_PROFIT_AFTER_TAX']),
        'TotalAssets':('balance',['BS_TOTAL_ASSETS']),
        'Equity':('balance',['BS_EQUITY','BS_OWNERS_EQUITY','BS_TOTAL_EQUITY']),
    }
    if metric not in metric_ids:
        return pd.DataFrame(columns=['Period','Value'])
    ds,ids=metric_ids[metric]
    candidates=[ROOT/'data'/'raw'/f'{t}_{ds}.csv']
    if ds=='income':
        candidates += [ROOT/'data'/'raw'/f'{t}_income_statement.csv', ROOT/'data'/'raw'/f'{t}_income.csv']
    p=next((x for x in candidates if x.exists()),None)
    if p is None:
        return pd.DataFrame(columns=['Period','Value'])
    try:
        q=pd.read_csv(p)
    except Exception:
        return pd.DataFrame(columns=['Period','Value'])
    if q.empty or not {'period','id','value'}.issubset(q.columns):
        return pd.DataFrame(columns=['Period','Value'])
    q=q[q['id'].astype(str).str.upper().isin([x.upper() for x in ids])].copy()
    q['Value']=pd.to_numeric(q['value'],errors='coerce')
    q=q.dropna(subset=['Value'])
    if q.empty:
        return pd.DataFrame(columns=['Period','Value'])
    def _canon(v):
        ss=str(v).upper().strip().replace('_','-').replace(' ','')
        m=re.fullmatch(r'(20\d{2})-?Q([1-4])',ss)
        return f"{m.group(1)}-Q{m.group(2)}" if m else None
    q['Period']=q['period'].map(_canon)
    q=q.dropna(subset=['Period'])
    if metric in {'ROE','ROA'}:
        q['Value']=q['Value'].map(lambda x: x/100.0 if abs(float(x))>1.5 else float(x))
    def _ord(v):
        m=re.fullmatch(r'(20\d{2})-Q([1-4])',str(v))
        return int(m.group(1))*4+int(m.group(2)) if m else -1
    q['_ord']=q['Period'].map(_ord)
    return q.sort_values('_ord').drop_duplicates('Period',keep='last')[['Period','Value']]

def _security_peer_raw_mean_history(ticker, metric, max_peers=10):
    """Mean peer history from the same raw semantic files used by the target."""
    try:
        pf=_securities_peer_frame(ticker,max_peers)
    except Exception:
        return pd.DataFrame(columns=['Period','IndustryMean'])
    if pf is None or not len(pf) or 'Ticker' not in pf.columns:
        return pd.DataFrame(columns=['Period','IndustryMean'])
    t=str(ticker).upper().strip()
    peer_tickers=[str(x).upper().strip() for x in pf['Ticker'].tolist() if str(x).upper().strip()!=t][:max_peers]
    frames=[]
    for x in peer_tickers:
        h=_security_raw_metric_history(x,metric)
        if len(h):
            h=h.copy(); h['Ticker']=x
            frames.append(h)
    if not frames:
        return pd.DataFrame(columns=['Period','IndustryMean'])
    allh=pd.concat(frames,ignore_index=True)
    return allh.groupby('Period',as_index=False)['Value'].mean().rename(columns={'Value':'IndustryMean'})

def metric_chart(ticker,metric,title,percent=False):
    from scripts.financial_views import history_comparison, BASIS_LABELS
    h=history_comparison(ticker,metric,globals().get('_history_basis','YTD'),globals().get('_history_years',5))
    fig=go.Figure()
    if len(h):
        fig.add_trace(go.Scatter(x=h.DisplayPeriod,y=h.Value,name=ticker,mode='lines+markers',connectgaps=False))
        fig.add_trace(go.Scatter(x=h.DisplayPeriod,y=h.PeerMean,name='TB peer cùng kỳ',mode='lines+markers',connectgaps=False,line=dict(dash='dash')))
    fig.update_layout(title=title+' · '+BASIS_LABELS[globals().get('_history_basis','YTD')],height=330,legend=dict(orientation='h',y=-.2),margin=dict(t=50,b=70))
    if percent:fig.update_yaxes(tickformat='.1%')
    return fig

def _latest_metric_period(ticker,metric):
    try:
        h=entity_history(ticker)
        z=h[h.Metric.astype(str).eq(metric)].copy()
        z['Date']=z.Period.map(period_date); z=z.dropna(subset=['Date']).sort_values('Date')
        return str(z.iloc[-1].Period) if len(z) else 'kỳ gần nhất'
    except Exception:
        return 'kỳ gần nhất'

def latest_peer_bar_chart(ticker,metric,title,percent=False,max_peers=10):
    from scripts.financial_views import peer_metric_frame
    q=peer_metric_frame(ticker,metric)
    fig=go.Figure()
    if len(q):fig.add_trace(go.Bar(x=q.Ticker,y=q.Value,marker_color=['#116b55' if t==ticker else '#9eb5c5' for t in q.Ticker]))
    fig.update_layout(title=title+' · cùng kỳ/cơ sở',height=330,margin=dict(t=55,b=45))
    if percent:fig.update_yaxes(tickformat='.1%')
    return fig

def _growth_from_hist(metric):
    try:x=pd.read_csv(DATA/'bank_history_long.csv')
    except:return pd.DataFrame(columns=['Ticker',metric+'_Growth'])
    x=x[x.Metric.astype(str).eq(metric)].copy(); x['Value']=pd.to_numeric(x.Value,errors='coerce'); x['Date']=x.Period.map(period_date); x=x.dropna(subset=['Value','Date']).sort_values(['Ticker','Date'])
    rows=[]
    for t,g in x.groupby('Ticker'):
        vals=g.Value.tolist(); gr=vals[-1]/vals[max(0,len(vals)-5)]-1 if len(vals)>=2 and vals[max(0,len(vals)-5)] else np.nan; rows.append({'Ticker':t,metric+'_Growth':gr})
    return pd.DataFrame(rows)

def bank_rating_result(ticker,overrides=None):
    return rate_three_methodologies(ticker,overrides)

u=universe()
st.sidebar.markdown('## PHÂN TÍCH DOANH NGHIỆP')
st.sidebar.caption('Phiên bản đang chạy: '+BUILD_TAG)
entity_filter=st.sidebar.selectbox('Nhóm doanh nghiệp',['Tất cả','Ngân hàng','Công ty chứng khoán','Công ty bảo hiểm','Doanh nghiệp phi tài chính'])
maptype={'Ngân hàng':'BANK','Công ty chứng khoán':'SECURITIES','Công ty bảo hiểm':'INSURANCE','Doanh nghiệp phi tài chính':'CORPORATE'}
if entity_filter=='Công ty bảo hiểm':
    _et=u['EntityType'].fillna('').astype(str).str.upper() if 'EntityType' in u else pd.Series('',index=u.index)
    _mt=u['Methodology'].fillna('').astype(str).str.upper() if 'Methodology' in u else pd.Series('',index=u.index)
    _sector=u['Sector'].fillna('').astype(str).str.lower() if 'Sector' in u else pd.Series('',index=u.index)
    z=u[_et.isin(['INSURANCE','INSURER','INSURANCE_NONLIFE','INSURANCE_LIFE']) | _mt.str.startswith('INSURANCE') | (_sector.str.contains('bảo hiểm|bao hiem|insurance',regex=True) & ~_sector.str.contains('môi giới|moi gioi|broker',regex=True))]
else:
    z=u if entity_filter=='Tất cả' else u[u.EntityType.eq(maptype[entity_filter])]
ticker_labels={r.Ticker:f'{r.Ticker} — {r.CompanyName}' for _,r in z.iterrows()}
selected=st.sidebar.selectbox('Mã doanh nghiệp',z.Ticker.tolist(),format_func=lambda x,_names=ticker_labels:_names.get(x,x)) if len(z) else None
presentation=st.sidebar.toggle('Chế độ trình bày',value=False)

if not selected: st.error('Không có doanh nghiệp trong bộ lọc.'); st.stop()
meta=get_company(selected); s=get_snapshot(selected); val=valuation(selected,s); peer=_peer_frame_safe(selected)
industry_name=industry_label(selected) if meta['EntityType']!='INSURANCE' else 'Nhóm tham chiếu bảo hiểm'
# Resolve the sector template once for the selected company.
# get_template() returns (template_key, template_dict).
sector_template_key, sector_template = get_template(meta.get('EntityType'), meta.get('Sector'))
st.sidebar.markdown('---'); st.sidebar.markdown('## TRẠNG THÁI')

def _clean_text(x, fallback='N/A'):
    if x is None:
        return fallback
    try:
        if pd.isna(x):
            return fallback
    except Exception:
        pass
    t=str(x).strip()
    return fallback if t in ('', 'nan', 'None', '<NA>') else t

_entity_type=_clean_text(meta.get('EntityType'))
_sector=_clean_text(meta.get('Sector'))
_peer_fallback=_clean_text(meta.get('PeerGroup'))
_methodology=_clean_text(meta.get('Methodology'))
_peer_label=_dynamic_peer_label_ui(selected, _peer_fallback)
if _entity_type=='INSURANCE':
    from scripts.insurance_peers import peer_universe
    _peer_label=f'{len(peer_universe(selected))} peer bảo hiểm; chỉ tiêu nghiệp vụ lọc cùng loại hình'

# Exchange/stock exchange is intentionally omitted from the UI.
st.sidebar.success(_entity_type)
st.sidebar.caption(
    f"Ngành: {_sector}\n\n"
    f"Nhóm so sánh: {_peer_label}\n\n"
    f"Phương pháp: {_methodology}"
)

st.markdown(
    '<div class="platform-title">NỀN TẢNG PHÂN TÍCH, ĐỊNH GIÁ, M&amp;A &amp; XẾP HẠNG TÍN NHIỆM DOANH NGHIỆP VIỆT NAM</div>'
    '<div class="platform-disclaimer">LƯU Ý: SẢN PHẨM CHỈ NHẰM MỤC ĐÍCH NGHIÊN CỨU, KHÔNG DÙNG CHO MỤC ĐÍCH KHUYẾN NGHỊ ĐẦU TƯ VÀ/HOẶC XẾP HẠNG TÍN NHIỆM - TÁC GIẢ: LÊ HOÀNG QUÂN - ĐT: 0384775999</div>',
    unsafe_allow_html=True
)
st.subheader(f"{selected} — {meta.get('CompanyName')}")

# Header KPIs adapt by entity.
if meta['EntityType']=='BANK': kpis=[('Giá thị trường',money(val.get('Price'))),('P/B',mult(s.get('PB'))),('ROE',pct(s.get('ROE'))),('NPL',pct(s.get('NPL'))),('CAR',pct(s.get('CAR'))),('CASA',pct(s.get('CASA')))]
elif meta['EntityType']=='SECURITIES': kpis=[('Giá thị trường',money(val.get('Price'))),('P/B',mult(s.get('PB'))),('P/E',mult(s.get('PE'))),('ROE',pct(s.get('ROE'))),('Nợ/VCSH',mult(s.get('DebtEquity'))),('Vốn khả dụng',pct((latest_capital_ratio(selected) or s.get('AvailableCapitalRatio'))))]
elif meta['EntityType']=='INSURANCE':kpis=[('Giá thị trường',money(val.get('Price'))),('Tổng tài sản',bn(s.get('TotalAssets'))),('Vốn chủ sở hữu',bn(s.get('Equity'))),('Dự phòng nghiệp vụ',bn(s.get('InsuranceReserves'))),('ROE',pct(s.get('ROE'))),('ROA',pct(s.get('ROA')))]
else:kpis=[('Giá thị trường',money(val.get('Price'))),('P/E',mult(s.get('PE'))),('ROE',pct(s.get('ROE'))),('ROA',pct(s.get('ROA'))),('Nợ/EBITDA',mult(s.get('DebtEBITDA'))),('Thanh toán hiện hành',mult(s.get('CurrentRatio')))]
cols=st.columns(6)
for c,(lab,v) in zip(cols,kpis):c.metric(lab,v)
if meta['EntityType']=='INSURANCE' and val.get('Price') is None:
    st.caption('Giá thị trường: thiếu dữ liệu giá đóng cửa trong nguồn hiện có; không suy ra từ P/E hoặc P/B.')

if meta.get('Methodology')=='EXCLUDED_SPECIALIZED' and not (str(meta.get('EntityType','')).upper().startswith('INSURANCE') or any(x in str(meta.get('Sector','')).lower() for x in ['bảo hiểm','insurance','bao hiem'])): st.warning('Ngành/loại hình này cần phương pháp XHTN chuyên biệt. App vẫn cho phép phân tích tài chính và định giá, nhưng không phát hành kết quả XHTN tự động.')
if meta['EntityType']!='BANK' and not any(num(s.get(k)) is not None for k in ['TotalAssets','Revenue','ROE','Price']): st.info(f'Chưa có đủ dữ liệu tài chính cho {selected}.')

_sections=['HỒ SƠ DOANH NGHIỆP','PHÂN TÍCH, ĐỊNH GIÁ & M&A','BÁO CÁO XẾP HẠNG TÍN NHIỆM','DỮ LIỆU & QUẢN TRỊ']
_active_section=st.radio('Điều hướng',_sections,horizontal=True,label_visibility='collapsed',key='main_section_nav')
from scripts.financial_views import BASIS_LABELS, HISTORY_METRICS, HISTORY_NOTE, history_table
_history_basis=st.sidebar.selectbox('Cơ sở lịch sử',list(BASIS_LABELS),index=2 if meta['EntityType']=='INSURANCE' else 0,format_func=lambda k:BASIS_LABELS[k])
_history_years=st.sidebar.selectbox('Số năm cùng kỳ',[5,4,3])
if _active_section in _sections[:3]:
    with st.expander('Lịch sử doanh nghiệp 3–5 năm cùng kỳ',expanded=True):
        st.dataframe(history_table(selected,_history_basis,_history_years),hide_index=True,use_container_width=True)
        st.caption(HISTORY_NOTE)
        _history_options=HISTORY_METRICS[meta['EntityType']]
        _hm=st.multiselect('Chỉ tiêu trên biểu đồ lịch sử',_history_options,default=_history_options[:2],key='history_metrics_'+selected)
        for _i in range(0,len(_hm),2):
            _hc=st.columns(2)
            for _j,_m in enumerate(_hm[_i:_i+2]):
                from scripts.methodology_kpi_engine import LABELS, PCT
                with _hc[_j]:_plotly_chart_safe(metric_chart(selected,_m,LABELS.get(_m,_m),_m in PCT),use_container_width=True)


def _safe_show(df, wanted, n=None):
    if df is None or not len(df): return pd.DataFrame()
    cols=[c for c in wanted if c in df.columns]
    q=df[cols].copy() if cols else df.copy()
    return q.head(n) if n else q

def _rating_report_fallback(ticker, rating_result, error):
    """Minimal transparent DOCX when legacy report generator rejects missing rating values."""
    from docx import Document
    from docx.shared import Pt
    doc=Document()
    doc.add_heading('BÁO CÁO MÔ PHỎNG XẾP HẠNG TÍN NHIỆM', 0)
    doc.add_paragraph(f'Doanh nghiệp: {ticker}')
    doc.add_paragraph('BẢN DỰ THẢO – KHÔNG PHẢI KẾT QUẢ XHTN ĐƯỢC PHÊ DUYỆT')
    r=rating_result or {}
    for k in ('MethodologyName','Anchor','SACP','SCA','ICR','Outlook'):
        v=r.get(k)
        if v is not None: doc.add_paragraph(f'{k}: {v}')
    doc.add_heading('Các cấu phần',level=1)
    for k,v in (r.get('Factors') or {}).items(): doc.add_paragraph(f'{k}: {v}')
    doc.add_heading('Bằng chứng và dữ liệu cần bổ sung',level=1)
    for k,v in (r.get('MissingEvidence') or {}).items() if isinstance(r.get('MissingEvidence'),dict) else []:
        doc.add_paragraph(f'{k}: {v}')
    doc.add_paragraph('Bộ tạo báo cáo chi tiết chưa xử lý được dữ liệu thiếu; bản này chỉ trình bày kết quả thực tế có sẵn, không suy diễn bậc XHTN.')
    out=io.BytesIO();doc.save(out);return out.getvalue()


def _download_report_block(report_type, rating_result=None):
    title='Báo cáo Phân tích – Định giá – M&A' if report_type=='analysis' else 'Báo cáo mô phỏng quá trình Xếp hạng tín nhiệm'
    st.markdown('---');st.subheader('Xuất '+title)
    c1,c2=st.columns(2)
    suffix='Phan_tich_Dinh_gia_MA' if report_type=='analysis' else 'XHTN'
    import hashlib
    from scripts.financial_evidence import signature
    _report_sources=[signature(p) for p in (ROOT/'data/raw').glob('*.csv')]+[signature(ROOT/'config'/p) for p in ['insurance_entities.csv','insurance_metric_inputs.csv','raw_period_basis.csv']] if meta['EntityType']=='INSURANCE' else []
    _digest=hashlib.sha256(json.dumps({'build':BUILD_TAG,'sources':_report_sources,'rating':rating_result,'basis':_history_basis,'years':_history_years,'snapshot':s.get('SnapshotPeriod')},sort_keys=True,default=str).encode()).hexdigest()[:12]
    identity=f'{selected}_{report_type}_{_digest}'
    for kind, column, title, mime, ext in [
        ('docx',c1,'Word','application/vnd.openxmlformats-officedocument.wordprocessingml.document','docx'),
        ('pdf',c2,'PDF','application/pdf','pdf')]:
        key=f'report_{identity}_{kind}'
        with column:
            if st.button('Chuẩn bị báo cáo '+title,key='prepare_'+key,use_container_width=True):
                try:
                    data=(generate_docx(selected,report_type,rating_result,{'history_basis':_history_basis,'history_years':_history_years}) if kind=='docx' else generate_pdf(selected,report_type,rating_result,{'history_basis':_history_basis,'history_years':_history_years}))
                    st.session_state[key]=data.getvalue() if hasattr(data,'getvalue') else data
                except Exception as exc:
                    if kind=='docx' and report_type=='rating':
                        st.session_state[key]=_rating_report_fallback(selected,rating_result,str(exc))
                        st.warning('Báo cáo chi tiết chưa tạo được; đã chuẩn bị bản Word dự thảo với dữ liệu thực có. Lỗi: '+str(exc))
                    else:
                        st.session_state.pop(key,None)
                        st.error('Không thể tạo báo cáo '+title+': '+str(exc))
            data=st.session_state.get(key)
            if isinstance(data,(bytes,bytearray)) and data:
                st.download_button('Tải báo cáo '+title,data,file_name=f'{selected}_{suffix}.{ext}',mime=mime,key='download_'+key,use_container_width=True)

if _active_section==_sections[0]:
    st.subheader('Thông tin tổng quan')
    _profile=load_company_profile(selected)
    if _profile:
        st.write(profile_overview_text(selected))
        _src=profile_source_text(selected)
        if _src:
            st.caption('Nguồn hồ sơ: '+_src)
    else:
        st.info('Chưa có hồ sơ doanh nghiệp được xác minh từ nguồn công khai. Không tự suy diễn phần tổng quan từ dữ liệu Vnstock.')

    st.subheader('Hồ sơ doanh nghiệp')
    a,b,c=st.columns(3)
    a.metric('Ngành',meta.get('Sector','N/A'))
    b.metric('Nhóm so sánh',industry_name)
    c.metric('Phương pháp',meta.get('Methodology','N/A'))
    st.write(f"**{meta.get('CompanyName')}** được hệ thống tự phân loại vào **{meta.get('Sector')}**. "
             f"Methodology XHTN: **{meta.get('Methodology')}**.")
    st.caption(f"Nguồn phân ngành: {meta.get('IndustrySource','Master/Legacy')} · Cấp ICB: {meta.get('IndustryLevelUsed','N/A')}")
    try:
        from scripts.dynamic_peer_engine import select_dynamic_peers
        _dp=select_dynamic_peers(selected)
        if len(_dp):
            with st.expander('Nhóm doanh nghiệp tương đồng được hệ thống tự chọn'):
                _show=[c for c in ['PeerRank','Ticker','CompanyName','Sector','SimilarityScore','SimilarityBasis','PeerPool'] if c in _dp.columns]
                _v=_dp[_show].copy()
                if 'SimilarityScore' in _v.columns:_v['SimilarityScore']=(_v['SimilarityScore']*100).round(1).astype(str)+'%'
                st.dataframe(_v,hide_index=True,use_container_width=True)
                st.caption('Nhóm tương đồng này phục vụ định giá; biểu đồ tài chính chung dùng toàn bộ nhóm bảo hiểm.' if meta['EntityType']=='INSURANCE' else 'Peer được tái tính sau mỗi Full Refresh từ toàn bộ universe Vnstock; Master chỉ còn là metadata/fallback.')
    except Exception:
        pass
    st.markdown('### Hồ sơ tài chính & xu hướng')
    if meta['EntityType']=='BANK':
        ml=[('TotalAssets','Tổng tài sản',False),('GrossLoans','Cho vay khách hàng',False),
            ('CustomerDeposits','Tiền gửi khách hàng',False),('ROE','ROE',True),('ROA','ROA',True),
            ('NIM','NIM',True),('NPL','Nợ xấu',True),('CAR','CAR',True),('CASA','CASA',True),('LDR','LDR',True)]
    elif meta['EntityType']=='INSURANCE':
        from scripts.insurance_analysis import render_analysis
        render_analysis(selected,_history_basis,_history_years,'profile_ins_')
        ml=[]
    else:
        ml=[('Revenue','Doanh thu',False),('NPAT','Lợi nhuận sau thuế',False),('ROE','ROE',True),
            ('ROA','ROA',True),('DebtEquity','Nợ/VCSH',False),('CurrentRatio','Thanh toán hiện hành',False)]
    for i in range(0,len(ml),2):
        cc=st.columns(2)
        for j,(m,t,pf) in enumerate(ml[i:i+2]):
            with cc[j]:
                _plotly_chart_safe(metric_chart(selected,m,f'{t} · xu hướng DN và trung bình peer',pf),use_container_width=True)
                _plotly_chart_safe(latest_peer_bar_chart(selected,m,f'{t} · DN và 10 peer tại kỳ gần nhất',pf),use_container_width=True)
    st.markdown('### Bộ chỉ tiêu theo methodology & nhóm tương đồng động')
    st.caption('Benchmark loại chính doanh nghiệp đang phân tích; chỉ hiển thị trung bình/trung vị khi có tối thiểu 5 peer có dữ liệu cho chỉ tiêu đó.')
    skpi,_,_=sector_kpi_table(selected)
    if meta.get('EntityType')=='SECURITIES':
        _peer_obs0=0
        try:
            _peer_obs0=int(pd.to_numeric(skpi.get('Số DN có dữ liệu'),errors='coerce').fillna(0).max()) if len(skpi) else 0
        except Exception:
            _peer_obs0=0
        if _peer_obs0==0 and skpi.empty:
            skpi=_securities_kpi_fallback(selected,peer)
    if len(skpi):
        wanted=['Nhóm phân tích','Chỉ tiêu','Doanh nghiệp','Trung bình ngành','Trung vị ngành',
                'Số DN có dữ liệu','Chênh lệch với TB ngành','Trạng thái dữ liệu','Kỳ dữ liệu','Cơ sở','Giải thích','Trạng thái peer']
        _kpi_display=_safe_show(skpi,wanted)
        if meta['EntityType']=='INSURANCE':
            _kpi_display=_kpi_display.rename(columns={'Trung bình ngành':'Bình quân nhóm tham chiếu','Trung vị ngành':'Trung vị nhóm tham chiếu','Chênh lệch với TB ngành':'Chênh lệch với bình quân nhóm'})
        st.dataframe(_kpi_display,hide_index=True,use_container_width=True)
    else: st.info('Chưa có dữ liệu methodology KPI.')
    if len(peer) and meta['EntityType']!='INSURANCE':
        st.markdown('### Vị trí tương đối trong peer group')
        q=peer.copy(); q['Doanh nghiệp']=q.Ticker.astype(str)
        cols_show=['Doanh nghiệp','ROE','ROA','PB','PE','DebtEquity','CurrentRatio','NPL','CAR','CASA','NIM','CIR','LDR']
        st.dataframe(_safe_show(q,cols_show),hide_index=True,use_container_width=True)

if _active_section==_sections[1]:
    st.subheader('Phân tích, Định giá & M&A')
    aa=intelligent_analyze(selected)
    c1,c2,c3=st.columns(3)
    c1.metric('Quan điểm định lượng',aa.get('View','N/A')); c2.metric('Điểm tín hiệu',aa.get('Score','N/A')); c3.metric('Mẫu chuyên ngành',aa.get('Template','N/A'))
    st.write(aa.get('Conclusion',''))
    l,r=st.columns(2)
    with l:
        st.markdown('#### Điểm mạnh tương đối')
        for x in aa.get('Strengths',[]): st.write('• '+x)
    with r:
        st.markdown('#### Rủi ro / điểm yếu tương đối')
        for x in aa.get('Risks',[]): st.write('• '+x)

    st.markdown('### Định giá')
    vr=valuation_regime(selected); fv=fair_value_range(selected); vt=triangulate(selected)
    if val.get('Status')!='OK':st.info(val.get('Reason','Thiếu dữ liệu định giá'))
    a,b,c,d=st.columns(4)
    a.metric('Giá thị trường',money(val.get('Price'))); b.metric('Giá trị tham chiếu',money(val.get('FairValue')))
    c.metric('Tiềm năng',pct(val.get('Upside'))); d.metric('Chế độ định giá',vr.get('Regime','N/A'))
    q1,q2,q3,q4=st.columns(4)
    fmt=lambda x:'N/A' if x is None else f"{x:,.0f} đồng/cp".replace(',','.')
    q1.metric('Bear',fmt(fv.get('Bear')))
    q2.metric('Base',fmt(fv.get('Base')))
    q3.metric('Bull',fmt(fv.get('Bull')))
    q4.metric('Chiến lược/M&A',fmt(fv.get('StrategicMA')))
    _sa=fv.get('ScenarioAssumptions',{})
    if _sa:
        st.caption(
            "Bear: "+str(_sa.get('Bear',''))+"  |  "
            "Base: "+str(_sa.get('Base',''))+"  |  "
            "Bull: "+str(_sa.get('Bull',''))
        )
    lenses=pd.DataFrame(vt.get('Lenses',[]))
    if len(lenses): st.dataframe(lenses,hide_index=True,use_container_width=True)

    st.markdown('### M&A, quyền kiểm soát, tái cấu trúc & kịch bản stress')
    base=num(val.get('FairValue'))
    c1,c2,c3=st.columns(3)
    premium=c1.slider('Thặng dư quyền kiểm soát',0.0,0.60,0.15,0.01)
    synergy=c2.slider('Giá trị cộng hưởng',0.0,0.50,0.08,0.01)
    stake=c3.slider('Tỷ lệ mua',0.01,1.0,0.51,0.01)
    strategic=base*(1+premium+synergy) if base else None
    st.metric('Giá trị chiến lược tham chiếu/cp',money(strategic))
    st.caption('Control premium, synergy và tỷ lệ mua là biến kịch bản theo từng thương vụ; không hard-code giả định của STB cho doanh nghiệp khác.')
    x1,x2=st.columns(2)
    with x1:
        debt_cut=st.slider('Giảm nợ giả định',0,50,10,5)
        equity_raise=st.slider('Tăng vốn giả định',0,50,10,5)
        st.write(f'Kịch bản tái cấu trúc: giảm nợ {debt_cut}% · tăng vốn {equity_raise}%.')
    with x2:
        if meta['EntityType']=='BANK':
            shock=st.slider('Shock NPL (điểm %)',0.0,5.0,1.0,.25)
            st.write(f'NPL hiện tại {pct(s.get("NPL"))}; stress cộng thêm {vi(shock,2)} điểm %.')
        else:
            rev=st.slider('Shock doanh thu',-50,20,-10,5); margin=st.slider('Shock biên lợi nhuận',-10,10,-2,1)
            st.write(f'Doanh thu {rev:+d}% · biên lợi nhuận {margin:+d} điểm %.')
    if meta['EntityType']=='INSURANCE':
        from scripts.insurance_analysis import render_analysis
        render_analysis(selected,_history_basis,_history_years,'valuation_ins_')
    _download_report_block('analysis')

if _active_section==_sections[2]:
    st.subheader('Báo cáo Xếp hạng tín nhiệm')
    _rating_overrides={}
    _entity_method=str(meta.get('EntityType','')).upper()
    _insurance_selected=(_entity_method.startswith('INSURANCE') or str(meta.get('Methodology','')).upper().startswith('INSURANCE') or any(x in str(meta.get('Sector','')).lower() for x in ['bảo hiểm','insurance','bao hiem']))
    if _insurance_selected:
        from scripts.insurance_ui import render_assessment
        _rating_overrides=render_assessment(selected)
    elif _entity_method=='BANK':
        with st.expander('Thẩm định yếu tố ngân hàng và hỗ trợ bên ngoài',expanded=False):
            st.caption('Anchor A− cố định. Chưa đủ chỉ tiêu và bằng chứng thì giữ N/A; không mặc định chấm trung tính.')
            _grades=['Tự tính từ dữ liệu','Rất Mạnh','Mạnh','Phù Hợp','Trung Bình','Yếu','Rất Yếu']
            for _key,_lab in [('BusinessProfile','Hồ sơ kinh doanh'),('CapitalEarnings','Vốn và lợi nhuận'),('RiskPosition','Vị thế rủi ro'),('FundingLiquidity','Huy động vốn và thanh khoản')]:
                _g=st.selectbox(_lab,_grades,key=f'bank_{selected}_{_key}')
                if _g!='Tự tính từ dữ liệu':_rating_overrides[_key]=_g
            _confirmed=st.checkbox('Đã đánh giá mức hỗ trợ bên ngoài',key=f'bank_{selected}_support_confirmed')
            if _confirmed:_rating_overrides['ExternalSupportNotches']=st.number_input('Notch hỗ trợ đã thẩm định',0,8,0,1,key=f'bank_{selected}_support')
    elif _entity_method=='SECURITIES':
        with st.expander('Thẩm định các yếu tố công ty chứng khoán',expanded=False):
            for _key,_lab in [('BusinessProfile','Hồ sơ kinh doanh'),('CapitalEarnings','Vốn và lợi nhuận'),('RiskPosition','Vị thế rủi ro'),('FundingLiquidity','Nguồn vốn và thanh khoản')]:
                _v=st.selectbox(_lab,['Chưa thẩm định','Rất Mạnh','Mạnh','Phù Hợp','Trung Bình','Yếu','Rất Yếu'],key=f'sec_{selected}_{_key}')
                if _v!='Chưa thẩm định':_rating_overrides[_key]=_v
            _v=st.selectbox('Notch hỗ trợ bên ngoài',['Chưa thẩm định']+list(range(9)),key=f'sec_{selected}_support')
            if isinstance(_v,int):_rating_overrides['ExternalSupportNotches']=_v
    else:
        with st.expander('Thẩm định yếu tố định tính doanh nghiệp phi tài chính',expanded=False):
            for _key,_lab in [('IndustryRisk','Rủi ro vĩ mô và ngành'),('GovernanceRisk','Rủi ro quản trị và quản lý')]:
                _v=st.selectbox(_lab,['Chưa thẩm định']+list(range(1,7)),key=f'corp_{selected}_{_key}')
                if isinstance(_v,int):_rating_overrides[_key]=_v
            for _key,_lab,_vals in [('ModifierNotches','Điều chỉnh khác',range(-6,7)),('ExternalSupportNotches','Hỗ trợ bên ngoài',range(9))]:
                _v=st.selectbox(_lab,['Chưa thẩm định']+list(_vals),key=f'corp_{selected}_{_key}')
                if isinstance(_v,int):_rating_overrides[_key]=_v
    rr3=rate_three_methodologies(selected,_rating_overrides)
    if rr3.get('Methodology')=='INSURANCE':
        from scripts.insurance_rating import insurance_scorecard, missing_evidence_labels, rating_status_text, rating_value_text
        from scripts.insurance_presentation import pending_assessment_groups
        if rr3.get('RatingStatus')=='SIMULATED':
            st.success('Đã tính xong kết quả mô phỏng từ các đầu vào đã nhập; cần phê duyệt theo quy trình XHTN.')
        elif rr3.get('RatingStatus')=='NOT_APPLICABLE':
            st.info(rating_status_text(rr3)+'. '+rr3.get('ApplicabilityReason',''))
        else:
            st.info('Phân tích tài chính đã tách khỏi bước phê duyệt XHTN. Chưa đủ hồ sơ để xác định bậc mô phỏng; xem các nhóm cần bổ sung bên dưới.')
        if rr3.get('RatingBlocker')!='SCOPE':
            with st.expander('Hồ sơ cần bổ sung để xác định bậc XHTN',expanded=False):
                st.caption(f"Đầu vào đã nhập: {rr3['AssessmentInputCount']}/{rr3['AssessmentInputTotal']}. Đây là tiến độ hồ sơ, không phải độ đầy đủ BCTC.")
                _pending=pending_assessment_groups(rr3)
                if _pending:st.dataframe(pd.DataFrame(_pending),hide_index=True,use_container_width=True)
                st.write('; '.join(missing_evidence_labels(rr3)))
    elif rr3.get('ICR') in (None,'N/A'):st.info('Chưa hình thành bậc XHTN mô phỏng: '+str(rr3.get('MissingEvidence','Chưa đủ bằng chứng')))
    r1,r2,r3=st.columns(3)
    if rr3.get('Methodology')=='INSURANCE':
        r1.metric('Anchor',rating_value_text(rr3,'Anchor'));r2.metric('SCA',rating_value_text(rr3,'SCA'));r3.metric('ICR',rating_value_text(rr3,'ICR'))
    else:
        r1.metric('Anchor',_rating_no_vn(rr3.get('Anchor','N/A')));r2.metric('SACP / SCA',_rating_no_vn(rr3.get('SACP',rr3.get('SCA','N/A'))));r3.metric('ICR',_rating_no_vn(rr3.get('ICR','N/A')))

    if meta['EntityType']=='INSURANCE':
        from scripts.insurance_analysis import render_analysis
        render_analysis(selected,_history_basis,_history_years,'rating_ins_')
    if meta['EntityType']!='INSURANCE':st.markdown('### Phân tích các chỉ số tín dụng trọng yếu')
    _rating_metric_map={
        'BANK':['ROE','ROA','NIM','NPL','CAR','CASA','LDR','CIR'],
        'SECURITIES':['ROE','ROA','AvailableCapitalRatio','DebtEquity','CurrentRatio','PB','PE'],
        'INSURANCE':['ROE','ROA','CombinedRatio','SolvencyRatio','LossRatio','ExpenseRatio','NetInsurancePremium'],
        'CORPORATE':['Revenue','GrossMargin','EBITDAMargin','ROE','ROA','DebtEquity','DebtEBITDA','CFO_Debt','FOCF_Debt','CurrentRatio','CashDebt']
    }
    _rkpi,_,_=sector_kpi_table(selected)
    if meta.get('EntityType')=='SECURITIES':
        _peer_obs=0
        try:
            _peer_obs=int(pd.to_numeric(_rkpi.get('Số DN có dữ liệu'),errors='coerce').fillna(0).max()) if len(_rkpi) else 0
        except Exception:
            _peer_obs=0
        if _peer_obs==0 and _rkpi.empty:
            _rkpi=_securities_kpi_fallback(selected,peer)
    if len(_rkpi) and meta.get('EntityType')!='INSURANCE':
        _want=set(_rating_metric_map.get(meta.get('EntityType'),[]))
        _show=_rkpi[_rkpi.Metric.astype(str).isin(_want)].copy() if 'Metric' in _rkpi.columns else _rkpi.copy()
        _cols=['Nhóm phân tích','Chỉ tiêu','Doanh nghiệp','Trung bình ngành','Trung vị ngành','Số DN có dữ liệu','Chênh lệch với TB ngành','Trạng thái dữ liệu','Kỳ dữ liệu','Cơ sở','Giải thích','Trạng thái peer']
        st.dataframe(_safe_show(_show,_cols),hide_index=True,use_container_width=True)
        # Current-period cross-sectional charts for the most decision-relevant metrics.
        if meta.get('EntityType')=='BANK': _bar_metrics=[('ROE','ROE',True),('NPL','Nợ xấu',True),('CAR','CAR',True),('CASA','CASA',True)]
        elif meta.get('EntityType')=='INSURANCE': _bar_metrics=[('ROE','ROE',True),('ROA','ROA',True),('EquityAssets','VCSH/TTS',True),('InsuranceReservesEquity','Dự phòng nghiệp vụ/VCSH',False)]
        elif meta.get('EntityType')=='SECURITIES': _bar_metrics=[('ROE','ROE',True),('DebtEquity','Nợ/VCSH',False),('CurrentRatio','Thanh toán hiện hành',False),('AvailableCapitalRatio','Vốn khả dụng',True)]
        else: _bar_metrics=[('ROE','ROE',True),('DebtEquity','Nợ/VCSH',False),('DebtEBITDA','Nợ/EBITDA',False),('CurrentRatio','Thanh toán hiện hành',False)]
        for _i in range(0,len(_bar_metrics),2):
            _cc=st.columns(2)
            for _j,(_m,_t,_pf) in enumerate(_bar_metrics[_i:_i+2]):
                with _cc[_j]: _plotly_chart_safe(latest_peer_bar_chart(selected,_m,f'{_t} · DN và 10 peer',_pf),use_container_width=True)
    elif meta.get('EntityType')!='INSURANCE':
        st.info('Chưa có đủ chỉ tiêu định lượng để phân tích XHTN.')

    if rr3.get('Methodology')=='CORPORATE':
        labels=rr3.get('RiskLabels',{}); _rows=[]
        for k,v in rr3.get('RiskScores',{}).items():
            try: _score=f'{float(v):.1f}/6' if pd.notna(v) else 'N/A'
            except Exception: _score='N/A'
            _rows.append({'Nhóm rủi ro':k,'Điểm':_score,'Mức rủi ro':labels.get(k,'N/A')})
        st.markdown('### Tổng hợp hồ sơ rủi ro')
        st.dataframe(pd.DataFrame(_rows),hide_index=True,use_container_width=True)
        p1,p2,p3=st.columns(3); p1.metric('Peer có dữ liệu',rr3.get('PeerCount',0)); p2.metric('Thanh khoản',rr3.get('Liquidity','N/A')); p3.metric('Đủ dữ liệu để tự động XHTN','CÓ' if rr3.get('DataSufficientForAutoRating') else 'CHƯA')
        if not rr3.get('DataSufficientForAutoRating'):
            st.warning('Chưa đủ dữ liệu doanh nghiệp/peer để phát hành bậc XHTN mô phỏng.')
    elif rr3.get('Methodology')=='INSURANCE':
        st.markdown('### Tổng kết cấu phần xếp hạng')
        _ins_rows=insurance_scorecard(rr3)
        st.dataframe(pd.DataFrame(_ins_rows),hide_index=True,use_container_width=True)
        st.caption(rr3.get('ScopeNote',''))
    elif rr3.get('Methodology') in ('BANK','SECURITIES'):
        st.markdown('### Tổng hợp hồ sơ rủi ro')
        st.dataframe(pd.DataFrame([{'Yếu tố':k,'Đánh giá':v} for k,v in rr3.get('Factors',{}).items()]),hide_index=True,use_container_width=True)

    # Component summary: missing assessments remain unknown, never converted to +0.
    if rr3.get('Methodology')!='INSURANCE':st.markdown('### Tổng kết cấu phần xếp hạng')
    if rr3.get('Methodology') in ('BANK','SECURITIES'):
        _desc_score={'Rất Mạnh':'1/6','Mạnh':'2/6','Phù Hợp':'3/6','Trung Bình':'4/6','Yếu':'5/6','Rất Yếu':'6/6'}
        _desc_notch={'Rất Mạnh':2,'Mạnh':1,'Phù Hợp':0,'Trung Bình':-1,'Yếu':-2,'Rất Yếu':-4}
        _sumrows=[]
        if rr3.get('Methodology')=='SECURITIES':
            _sumrows.append({'Cấu phần':'BICRA tham chiếu','Đánh giá / Điểm':_rating_no_vn(rr3.get('BICRAReference','N/A')),'Nâng/Hạ notch':'—','Kết quả':'Tham chiếu'})
            _sumrows.append({'Cấu phần':'Điều chỉnh Anchor CTCK','Đánh giá / Điểm':'Đặc thù ngành CTCK','Nâng/Hạ notch':str(rr3.get('SectorAnchorAdjustment','N/A')),'Kết quả':_rating_no_vn(rr3.get('Anchor','N/A'))})
        else:
            _sumrows.append({'Cấu phần':'Điểm ban đầu ngành / BICRA','Đánh giá / Điểm':_rating_no_vn(rr3.get('BICRA',rr3.get('Anchor','N/A'))),'Nâng/Hạ notch':'—','Kết quả':_rating_no_vn(rr3.get('Anchor','N/A'))})
        _factors=rr3.get('Factors') or {}
        _factor_labels={'Hồ sơ Kinh doanh','Vốn và Lợi nhuận','Vị thế Rủi ro','Huy động vốn và Thanh khoản'} if rr3.get('Methodology')=='BANK' else set(_factors)
        for _k in _factor_labels:
            _v=_factors.get(_k)
            _sumrows.append({'Cấu phần':_k,'Đánh giá / Điểm':f"{_v} ({_desc_score[_v]})" if _v in _desc_score else 'Chưa thẩm định', 'Nâng/Hạ notch':f"{_desc_notch[_v]:+d}" if _v in _desc_notch else 'N/A','Kết quả':'Điều chỉnh nội sinh' if _v in _desc_notch else 'Thiếu đánh giá'})
        def _fmt_notch(v):
            try:return f'{int(v):+d}' if v is not None and pd.notna(v) else 'N/A'
            except (ValueError,TypeError):return 'N/A'
        _sumrows.append({'Cấu phần':'Tổng điều chỉnh nội sinh','Đánh giá / Điểm':'—','Nâng/Hạ notch':_fmt_notch(rr3.get('InternalNotches')),'Kết quả':_rating_no_vn(rr3.get('SACP','N/A'))})
        _support=rr3.get('ExternalSupportNotches')
        _sumrows.append({'Cấu phần':'Hỗ trợ bên ngoài','Đánh giá / Điểm':'Chưa thẩm định' if _support is None else ('Trung lập' if _support==0 else 'Có điều chỉnh'),'Nâng/Hạ notch':_fmt_notch(_support),'Kết quả':_rating_no_vn(rr3.get('ICR','N/A'))})
        _sumrows.append({'Cấu phần':'Kết quả XHTN','Đánh giá / Điểm':rr3.get('Outlook','Chờ thẩm định'),'Nâng/Hạ notch':'—','Kết quả':_rating_no_vn(rr3.get('ICR','N/A'))})
        st.dataframe(pd.DataFrame(_sumrows),hide_index=True,use_container_width=True)
    elif rr3.get('Methodology')=='CORPORATE':
        _rs=rr3.get('RiskScores',{}); _rl=rr3.get('RiskLabels',{})
        _sumrows=[]
        for _k,_v in _rs.items():
            try:_pt=f"{float(_v):.1f}/6"
            except Exception:_pt='N/A'
            _sumrows.append({'Cấu phần':_k,'Điểm':_pt,'Mức rủi ro':_rl.get(_k,'N/A'),'Tác động':'Điểm cấu phần'})
        _weighted=rr3.get('WeightedScore')
        try:_weighted_txt=f"{float(_weighted):.2f}/6"
        except Exception:_weighted_txt='N/A'
        def _corp_anchor_band_ui(x):
            from scripts.three_methodology_rating import _anchor_from_weighted
            return 'Chưa thẩm định' if num(x) is None else _anchor_from_weighted(float(x)).upper()
        def _corp_anchor_rating_ui(x):return _corp_anchor_band_ui(x)
        _sumrows += [
            {'Cấu phần':'Điểm rủi ro tổng hợp có trọng số','Điểm':_weighted_txt,'Mức rủi ro':'Điểm cuối trước modifier','Tác động':_corp_anchor_band_ui(_weighted)},
            {'Cấu phần':'Anchor quy đổi từ điểm tổng hợp','Điểm':_corp_anchor_rating_ui(_weighted),'Mức rủi ro':'Quy đổi theo engine mô phỏng','Tác động':'Điểm → Anchor'},
            {'Cấu phần':'Thanh khoản','Điểm':str(rr3.get('LiquidityScore','N/A')),'Mức rủi ro':rr3.get('Liquidity','N/A'),'Tác động':'Modifier/cap nếu trọng yếu; không cộng trực tiếp vào điểm tổng hợp'},
            {'Cấu phần':'Modifier','Điểm':'—','Mức rủi ro':'—','Tác động':('Chưa thẩm định' if rr3.get('ModifierNotches') is None else f"{int(rr3['ModifierNotches']):+d} notch")},
            {'Cấu phần':'SCA','Điểm':_rating_no_vn(rr3.get('SCA','N/A')),'Mức rủi ro':'Sau modifier/cap','Tác động':'Anchor + modifier/cap'},
            {'Cấu phần':'Hỗ trợ bên ngoài','Điểm':'—','Mức rủi ro':'—','Tác động':('Chưa thẩm định' if rr3.get('ExternalSupportNotches') is None else f"{int(rr3['ExternalSupportNotches']):+d} notch")},
            {'Cấu phần':'Kết quả XHTN','Điểm':_rating_no_vn(rr3.get('ICR','N/A')),'Mức rủi ro':rr3.get('Outlook','N/A'),'Tác động':'SCA + hỗ trợ bên ngoài'},
        ]
        st.dataframe(pd.DataFrame(_sumrows),hide_index=True,use_container_width=True)

    # Show an auditable factor waterfall even before the final rating is approved.
    if rr3.get('Methodology')=='BANK':
        _wf=[{'Bước':'Anchor cố định','Điều chỉnh':'—','Bậc XHTN':_rating_no_vn(rr3.get('Anchor','N/A'))}]
        for _k,_v in (rr3.get('Factors') or {}).items():
            _wf.append({'Bước':_k,'Điều chỉnh':f"{_desc_notch[_v]:+d}" if _v in _desc_notch else 'N/A','Bậc XHTN':'Đang tổng hợp'})
        _wf.extend([{'Bước':'SACP','Điều chỉnh':'—','Bậc XHTN':_rating_no_vn(rr3.get('SACP','N/A'))},
                    {'Bước':'Hỗ trợ bên ngoài','Điều chỉnh':_fmt_notch(rr3.get('ExternalSupportNotches')),'Bậc XHTN':_rating_no_vn(rr3.get('ICR','N/A'))}])
        rc={'Waterfall':_wf}
    else:
        rc=committee_pack(selected,_rating_overrides)
    if rr3.get('Methodology')!='INSURANCE':
        st.markdown('### Waterfall trình Hội đồng XHTN')
        _waterfall=pd.DataFrame(rc.get('Waterfall',[]))
        if len(_waterfall):st.dataframe(_waterfall,hide_index=True,use_container_width=True)
        else:st.info('Chưa có waterfall; cần hoàn thành thẩm định cấu phần.')
    if rr3.get('Methodology')=='BANK':
        _spec={'CAR','ROA','ROE','NIM','CIR','NPL','LoanLossCoverage','CreditCost','Group2LoansRatio','CASA','LDR','LiquidityCoverageRatio','StableFundingRatio'}
        _present={m for m in _spec if num(s.get(m)) is not None}
        _cov=len(_present)/len(_spec)
        ev={'ICR':rr3.get('ICR','N/A'),'RatingConfidence':'Chờ thẩm định' if rr3.get('ICR') in ('N/A',None) else 'Mô phỏng – cần phê duyệt','DataQuality':{'Coverage':_cov},'EvidenceLedger':[]}
        st.caption(f'Coverage định lượng: {len(_present)}/{len(_spec)} chỉ tiêu thiết kế; không bao gồm đánh giá định tính và hỗ trợ bên ngoài.')
    else:
        ev=rating_evidence(selected,_rating_overrides)
        if rr3.get('Methodology')=='INSURANCE':
            _dq=ev['DataQuality']
            st.markdown('### Độ đầy đủ dữ liệu phân tích')
            e1,e2,e3=st.columns(3)
            e1.metric('Có số / chỉ tiêu áp dụng',f"{len(_dq['Present'])}/{_dq['ApplicableTotal']}")
            e2.metric('Cần bổ sung',len(_dq['Missing']));e3.metric('Không áp dụng',len(_dq['NotApplicable']))
            st.caption(_dq['Note'])
    if rr3.get('Methodology')!='INSURANCE':
        st.markdown('### Mức độ tin cậy của kết quả')
        _coverage=ev.get('DataQuality',{}).get('Coverage')
        _coverage_label=f'{_coverage*100:.0f}%' if isinstance(_coverage,(int,float)) and np.isfinite(_coverage) else 'N/A'
        e1,e2,e3=st.columns(3);e1.metric('ICR mô phỏng',_rating_no_vn(ev.get('ICR','N/A')));e2.metric('Độ tin cậy XHTN',ev.get('RatingConfidence','N/A'));e3.metric('Độ đầy đủ dữ liệu',_coverage_label)

    with st.expander('Phương pháp, audit trail & chi tiết máy tính',expanded=False):
        st.write(f"**Phương pháp tự động lựa chọn:** {rr3.get('MethodologyName','N/A')}")
        if rr3.get('Audit'): st.caption(rr3.get('Audit',''))
        if rr3.get('EvidenceMetrics'): st.caption('Chỉ tiêu định lượng sử dụng: '+', '.join(rr3.get('EvidenceMetrics',[])))
        ledger=pd.DataFrame(ev.get('EvidenceLedger',[]))
        if len(ledger): st.dataframe(ledger,hide_index=True,use_container_width=True)
        st.json(rr3)
    st.session_state['rating_result']=rr3
    _download_report_block('rating',rr3)

if _active_section==_sections[3]:
    st.subheader('Dữ liệu & Quản trị')
    from scripts.financial_evidence import STATUS
    _evidence_rows=[dict(r,Trạng_thái=STATUS.get(r['Status'],r['Status'])) for r in s.get('_MetricEvidence',{}).values()]
    st.markdown('### Bằng chứng dữ liệu của doanh nghiệp đang chọn')
    st.dataframe(pd.DataFrame(_evidence_rows),hide_index=True,use_container_width=True)
    if s.get('_DataIssues'):st.warning('; '.join(s['_DataIssues']))
    st.caption('Các bảng coverage và peer batch bên dưới là kết quả lần rebuild lưu gần nhất; bảng bằng chứng trên được đọc lại từ raw hiện tại.')
    try: cov=pd.read_csv(DATA/'coverage_matrix.csv')
    except Exception: cov=build_coverage_matrix() if build_coverage_matrix else pd.DataFrame()
    if len(cov):
        a,b,c,d=st.columns(4)
        a.metric('Universe trong Master',f'{len(cov):,}'.replace(',', '.'))
        a1=int((cov.Readiness=='PRODUCTION_READY').sum()) if 'Readiness' in cov else 0
        a2=int((cov.Readiness=='PARTIAL').sum()) if 'Readiness' in cov else 0
        a3=int((cov.Readiness=='INSUFFICIENT_DATA').sum()) if 'Readiness' in cov else 0
        b.metric('Sẵn sàng phân tích',f'{a1:,}'.replace(',', '.'));c.metric('Phủ một phần',f'{a2:,}'.replace(',', '.'));d.metric('Thiếu dữ liệu',f'{a3:,}'.replace(',', '.'))
        st.dataframe(cov,hide_index=True,use_container_width=True)
    st.markdown('### Dynamic peer engine')
    try:
        dps=pd.read_csv(DATA/'dynamic_peer_summary.csv')
        if len(dps):
            st.dataframe(dps,hide_index=True,use_container_width=True)
        else: st.info('Chưa có dữ liệu nhóm doanh nghiệp tương đồng.')
    except Exception:
        st.info('Dynamic peer map sẽ được tạo sau RUN_FULL_REFRESH.bat.')
    st.markdown('### Universe & methodology router')
    cols=[c for c in ['Ticker','CompanyName','EntityType','Sector','Exchange','PeerGroup','Methodology'] if c in u.columns]
    st.dataframe(u[cols],hide_index=True,use_container_width=True)

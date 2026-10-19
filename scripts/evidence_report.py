"""Word dùng cùng kết quả dữ liệu, peer và XHTN với giao diện."""
from io import BytesIO
from datetime import date
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from scripts.universal_data import get_company,get_snapshot,num
from scripts.financial_evidence import STATUS,SCOPE_NOTE,metric_evidence
from scripts.financial_views import history_table,history_comparison,HISTORY_METRICS,HISTORY_NOTE,BASIS_LABELS,fmt
from scripts.methodology_kpi_engine import methodology_kpi_table,LABELS,PCT,MULT
from scripts.multisector_valuation import valuation
from scripts.fair_value_range import fair_value_range
from scripts.three_methodology_rating import rate_company
from scripts.company_profile import profile_overview_text,profile_source_text

def _para(doc,text):
    p=doc.add_paragraph(str(text));p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def _heading(doc,text):doc.add_heading(text,1)

def _table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Mm(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            if widths:c.width=Mm(widths[j])
            tc=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            tc.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDE7EC' if i==0 else 'F5F7F8' if i%2==0 else 'FFFFFF');tc.append(sh)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
            tc.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(8.5);r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def _chart(ticker,metric,basis,years):
    h=history_comparison(ticker,metric,basis,years)
    if h.empty or h.Value.notna().sum()==0:return None
    scale=100 if metric in PCT else 1 if metric in MULT else 1e-9
    unit='%' if metric in PCT else 'lần' if metric in MULT else 'tỷ đồng'
    fig,ax=plt.subplots(figsize=(7.2,2.8),dpi=160)
    xs=np.arange(len(h));ax.plot(xs,pd.to_numeric(h.Value,errors='coerce')*scale,'o-',color='#166b55',label=ticker,lw=2)
    if h.PeerMean.notna().any():ax.plot(xs,h.PeerMean*scale,'s--',color='#607d93',label='TB peer cùng kỳ',lw=1.5)
    ax.set_xticks(xs,h.DisplayPeriod,fontsize=8);ax.set_ylabel(unit,fontsize=8)
    ax.set_title(LABELS.get(metric,metric),loc='left',fontsize=11,fontweight='bold');ax.tick_params(axis='y',labelsize=8)
    ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',alpha=.15);ax.legend(loc='best',fontsize=8,frameon=False)
    fig.tight_layout();bio=BytesIO();fig.savefig(bio,format='png',bbox_inches='tight');plt.close(fig);bio.seek(0);return bio

def _show_rating(v):
    return 'Chưa thẩm định' if v is None or str(v) in ('N/A','nan','None') else str(v).replace('vn','')

def generate_docx(ticker,report_type='analysis',rating_result=None,mna=None):
    ticker=str(ticker).upper();meta=get_company(ticker);s=get_snapshot(ticker);typ=meta['EntityType']
    if typ=='INSURANCE':
        from scripts.insurance_report import generate_insurance_docx
        return generate_insurance_docx(ticker,report_type,rating_result,mna)
    opts=mna or {};basis=opts.get('history_basis','YTD');years=int(opts.get('history_years',5))
    k=methodology_kpi_table(ticker);v=valuation(ticker,s);rr=rating_result if rating_result is not None else rate_company(ticker)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=Mm(15);sec.bottom_margin=Mm(15);sec.left_margin=Mm(16);sec.right_margin=Mm(16)
    normal=doc.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.08
    for name in ['Title','Subtitle','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='DejaVu Sans';style.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Title'].font.size=Pt(20);doc.styles['Heading 1'].font.size=Pt(13)
    title='Báo cáo mô phỏng xếp hạng tín nhiệm' if report_type=='rating' else 'Báo cáo phân tích định giá và M&A'
    doc.add_paragraph(title,'Title')
    doc.add_paragraph(f"{ticker}   {meta.get('DisplayName',meta.get('CompanyName',ticker))}",'Subtitle')
    _para(doc,f"Loại hình {typ} | Kỳ tài chính {s.get('SnapshotPeriod','Chưa có')} | Lịch sử {BASIS_LABELS[basis].lower()} {years} năm | Ngày lập {date.today():%d/%m/%Y}")
    _para(doc,'Tác giả Lê Hoàng Quân. Tài liệu nghiên cứu và mô phỏng; không phải kết quả xếp hạng đã được phê duyệt hoặc khuyến nghị đầu tư.')
    _heading(doc,'Kết quả và phạm vi phân tích')
    if report_type=='rating':
        _para(doc,f"Phương pháp: {rr.get('MethodologyName',typ)}. Anchor {_show_rating(rr.get('Anchor'))}; SACP/SCA {_show_rating(rr.get('SACP',rr.get('SCA')))}; ICR {_show_rating(rr.get('ICR'))}.")
        _para(doc,rr.get('Audit','Cần xác nhận bằng chứng và phạm vi phương pháp.'))
        if rr.get('RatingStatus')=='NOT_APPLICABLE':_para(doc,'Ma trận không quy định bậc cho tổ hợp điểm đã chọn; cần rà soát đầu vào và phạm vi phương pháp.')
        elif rr.get('ICR') in (None,'N/A'):_para(doc,'Chưa hình thành bậc xếp hạng mô phỏng do còn thiếu dữ liệu hoặc đầu vào thẩm định. Các chỉ tiêu định lượng dưới đây vẫn được trình bày theo dữ liệu thực có.')
    else:
        _para(doc,f"Phương pháp định giá: {v.get('PrimaryMethod')}. Giá thị trường {fmt('Price',v.get('Price'))} đồng/cp; giá trị tham chiếu {fmt('FairValue',v.get('FairValue'),v.get('Status'))}.")
        _para(doc,v.get('Reason',''))
    if typ=='INSURANCE':_para(doc,SCOPE_NOTE)
    try:
        profile=profile_overview_text(ticker)
        if profile:_heading(doc,'Hồ sơ doanh nghiệp');_para(doc,profile);_para(doc,profile_source_text(ticker))
    except (ValueError,KeyError,FileNotFoundError):pass
    _para(doc,f"Có {int(k['Doanh nghiệp'].notna().sum())}/{len(k)} chỉ tiêu định lượng trong ma trận phương pháp. Số liệu có sẵn không đồng nghĩa các yếu tố định tính đã được thẩm định.")

    _heading(doc,'Lịch sử tài chính cùng kỳ')
    h=history_table(ticker,basis,years)
    if len(h):
        widths=[43,15]+[120/(len(h.columns)-2)]*(len(h.columns)-2)
        _table(doc,list(h.columns),h.values.tolist(),widths)
    else:_para(doc,'Thiếu dữ liệu lịch sử có kỳ xác định.')
    _para(doc,HISTORY_NOTE)
    core=HISTORY_METRICS[typ]
    # Diễn giải trực tiếp thay đổi của doanh nghiệp, chỉ khi đủ hai mốc cùng kỳ.
    for metric in [core[0],'NPAT','ROE']:
        z=history_comparison(ticker,metric,basis,years)
        if len(z)>=2 and num(z.iloc[-1].Value) is not None and num(z.iloc[-2].Value) is not None:
            a,b=z.iloc[-2],z.iloc[-1];delta=b.Value-a.Value
            change=f'{delta*100:+.2f} điểm phần trăm' if metric in PCT else f'{delta/abs(a.Value)*100:+.1f}%' if a.Value!=0 else 'không tính tăng trưởng do kỳ gốc bằng 0'
            _para(doc,f"{LABELS.get(metric,metric)}: {fmt(metric,b.Value)} tại {b.DisplayPeriod}, so với {fmt(metric,a.Value)} tại {a.DisplayPeriod}; thay đổi {change}. Dữ liệu phản ánh chuỗi được cung cấp, cần đối chiếu khi phạm vi hợp nhất thay đổi.")
    chart_metrics={'BANK':['GrossLoans','NPL','ROE','CASA'],'CORPORATE':['Revenue','NPAT','DebtEquity','DebtEBITDA'],
                   'SECURITIES':['Revenue','NPAT','DebtEquity','ROE'],'INSURANCE':['NetInsurancePremium','NPAT','InsuranceReserves','ROE']}[typ]
    for i,m in enumerate(chart_metrics):
        if i%2==0:_heading(doc,'Diễn biến doanh nghiệp và nhóm so sánh')
        bio=_chart(ticker,m,basis,years)
        if bio:doc.add_picture(bio,width=Mm(174))
        else:_para(doc,LABELS.get(m,m)+': thiếu chuỗi cùng kỳ; không nội suy.')
    _heading(doc,'Chỉ tiêu hiện tại và đối chiếu peer')
    _para(doc,'Snapshot dùng số dư cuối kỳ và dòng tiền TTM. Bình quân và trung vị chỉ hiển thị khi có tối thiểu 5 peer cùng kỳ và cùng cơ sở; mẫu không bao gồm chính doanh nghiệp.')
    for group,q in k.groupby('Nhóm phân tích',sort=False):
        doc.add_heading(group,2)
        rows=[]
        for r in q.to_dict('records'):
            m=r['Metric'];rows.append([r['Chỉ tiêu'],fmt(m,r['Doanh nghiệp'],r['StatusCode']),
                fmt(m,r['TB ngành']),str(r['Số DN']),str(r['Kỳ dữ liệu']).replace('|',' / ')])
        _table(doc,['Chỉ tiêu','Doanh nghiệp','TB peer','N','Kỳ đầu vào'],rows,[43,33,31,9,62])
    _heading(doc,'Nguồn dữ liệu và các điểm cần bổ sung')
    missing=k[k.StatusCode.ne('OK')]
    if len(missing):
        _table(doc,['Chỉ tiêu','Trạng thái','Lý do'],[[r['Chỉ tiêu'],r['Trạng thái dữ liệu'],r['Giải thích']] for r in missing.to_dict('records')],[48,35,95])
    else:_para(doc,'Các chỉ tiêu trong ma trận hiện có đủ đầu vào định lượng. Việc xác minh chứng từ và các đánh giá định tính vẫn cần được hoàn thành riêng.')
    _para(doc,'Nguồn: CSV raw trong project và dữ liệu lịch sử ngân hàng lưu sẵn. Mã dòng được đối chiếu chính xác; các dòng order/level không phải giá trị tài chính. Bảng kỳ dài không cộng tỷ số hoặc số dư. Không dùng phép lấy thị giá chia P/E hoặc P/B khác kỳ để lấp dữ liệu EPS/BVPS.')
    evidence=[]
    for m in core[:7]:
        e=metric_evidence(s,m);evidence.append([LABELS.get(m,m),e['Basis'],e['Source'] or e['Reason']])
    _table(doc,['Chỉ tiêu','Cơ sở','Nguồn hoặc công thức'],evidence,[42,33,103])
    if typ!='INSURANCE':_para(doc,'EBIT = lợi nhuận trước thuế + chi phí lãi vay tuyệt đối; EBITDA = EBIT + khấu hao cùng cửa sổ TTM. Nợ vay là tổng vay ngắn hạn và dài hạn khi đủ cả hai dòng. Tổng các khoản cho vay CTCK không được đổi tên thành dư nợ ký quỹ.')
    _para(doc,'EPS suy tính dùng lợi nhuận cổ đông công ty mẹ TTM chia số cổ phiếu cuối kỳ; đây là proxy, không phải EPS cơ bản theo bình quân gia quyền. BVPS dùng vốn thuộc cổ đông công ty mẹ khi có số liệu lợi ích không kiểm soát.')
    if report_type=='analysis':
        _heading(doc,'Kịch bản định giá và M&A')
        fv=fair_value_range(ticker)
        _table(doc,['Kịch bản','Giá trị đồng mỗi cổ phiếu'],[[key,fmt('FairValue',fv.get(key),v.get('Status'))] for key in ['Bear','Base','Bull','StrategicMA']],[70,108])
        for key,txt in fv.get('ScenarioAssumptions',{}).items():_para(doc,key+': '+str(txt))
        _para(doc,'Thặng dư kiểm soát và cộng hưởng là giả định kịch bản; không phải dữ liệu quan sát. Thiếu giá trị cơ sở thì không tạo vùng giá hoặc giá trị M&A từ thị giá.')
    else:
        # Scorecard là phần cuối, sử dụng chính kết quả đã truyền từ giao diện.
        _heading(doc,'Tổng hợp cấu phần xếp hạng')
        rows=[['Anchor',_show_rating(rr.get('Anchor'))]]
        rows += [[k,_show_rating(v)] for k,v in rr.get('Factors',{}).items()]
        for key,value in rr.get('RiskScores',{}).items():rows.append([key,f'{value:.2f}/6' if num(value) is not None else 'Chưa thẩm định'])
        for key in ['CombinedRiskScore','FinancialRiskScore','WeightedScore','InternalNotches','ModifierNotches','ExternalSupportNotches']:
            if key in rr:rows.append([key,_show_rating(rr[key])])
        rows += [['SACP / SCA',_show_rating(rr.get('SACP',rr.get('SCA')))],['ICR',_show_rating(rr.get('ICR'))],['Triển vọng',rr.get('Outlook','Chưa thẩm định')]]
        if typ=='INSURANCE':
            from scripts.insurance_rating import insurance_scorecard, missing_evidence_labels
            rows=[[r['Cấu phần'],r['Kết quả']] for r in insurance_scorecard(rr)]
            from scripts.data_quality_engine import assess_data_quality
            dq=assess_data_quality(ticker)
            _para(doc,f"Dữ liệu định lượng cốt lõi: {len(dq['Present'])}/{len(dq['Present'])+len(dq['Missing'])} chỉ tiêu. Đầu vào thẩm định đã nhập: {rr.get('AssessmentInputCount',0)}/{rr.get('AssessmentInputTotal',9)}.")
            _para(doc,'Đã tính xong kết quả mô phỏng; chờ phê duyệt.' if rr.get('RatingStatus')=='SIMULATED' else 'Tổ hợp điểm không có bậc quy định trong ma trận.' if rr.get('RatingStatus')=='NOT_APPLICABLE' else 'Chưa hoàn thành đầu vào thẩm định; không tự gán bậc tín nhiệm.')
        _table(doc,['Cấu phần','Kết quả'],rows,[95,83])
        absent=rr.get('MissingEvidence',{})
        if typ=='INSURANCE':absent=missing_evidence_labels(rr)
        items=absent.items() if isinstance(absent,dict) else [(x,'Chưa thẩm định') for x in absent]
        for key,value in items:
            if isinstance(value,list):
                readable='; '.join(str(x.get('Metric',''))+': '+str(x.get('Reason','Thiếu dữ liệu')) if isinstance(x,dict) else str(x) for x in value)
            else:readable=str(value)
            _para(doc,f'{key}: {readable}')
        if typ=='BANK':_para(doc,'Anchor A− cố định. Thiếu đánh giá hồ sơ kinh doanh hoặc hỗ trợ bên ngoài không được xem là trung tính hoặc hỗ trợ 0 bậc.')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(ticker+' | ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for p in doc.paragraphs:p.paragraph_format.widow_control=True
    # Xóa border/theme màu của template mặc định, kể cả đường kẻ dưới Title.
    for el in list(doc.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
    for el in list(doc.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
    for style in [doc.styles[n] for n in ['Title','Subtitle','Heading 1','Heading 2']]:
        for color in style.element.iter(qn('w:color')):
            color.set(qn('w:val'),'000000')
            for attr in ['themeColor','themeShade','themeTint']:color.attrib.pop(qn('w:'+attr),None)
    out=BytesIO();doc.save(out);return out.getvalue()

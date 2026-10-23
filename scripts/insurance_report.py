"""Báo cáo bảo hiểm theo bố cục chung TCB/VDS/HPG, số liệu dùng chung UI."""
from io import BytesIO
from datetime import date
import re
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from scripts.insurance_analysis import analysis_pack, peer_universe, BASIS_TEXT
from scripts.insurance_peers import included_rows, POLICY_NOTE, SUBTYPE_LABELS
from scripts.insurance_rating import insurance_rating, insurance_scorecard, rating_status_text, rating_value_text
from scripts.insurance_presentation import quantitative_coverage, executive_findings, has_evidence, pending_assessment_groups,format_metric as fmt,data_gap_explanation
from scripts.financial_views import BASIS_LABELS
from scripts.financial_evidence import STATUS, insurance_reporting_subtype
from scripts.methodology_kpi_engine import PCT, MULT
from scripts.universal_data import get_company
from scripts.insurance_narrative import section_findings, sensitivity_narrative, overall_conclusion, ENGINE_VERSION

SECTIONS = [
    ('HỒ SƠ KINH DOANH', ['GrossWrittenPremium','NetInsurancePremium','Revenue','TotalAssets','MarketSharePremium','LargestProductShare']),
    ('HIỆU QUẢ KHAI THÁC VÀ KHẢ NĂNG SINH LỢI', ['NPAT','ROE','ROA','NetMargin','LossRatio','ExpenseRatio','CombinedRatio','UnderwritingProfit','InvestmentIncome']),
    ('VỐN VÀ DỰ PHÒNG NGHIỆP VỤ', ['Equity','EquityAssets','PaidInCapital','PaidInCapitalRequired','InsuranceReserves','InsuranceReservesEquity','ReservesPremium']),
    ('VỊ THẾ RỦI RO VÀ LINH HOẠT TÀI CHÍNH', ['FXMismatchLiabilities','InvestmentLeverage','LargestSectorShare','DebtEquity','DebtDue24MEquity','InsuranceEBITDAInterest']),
    ('THANH KHOẢN VÀ KHẢ NĂNG THANH TOÁN', ['StressedLiquidAssets','InsuranceContractObligations','InsuranceLiquidity','SolvencyRatio']),
    ('ĐỊNH GIÁ THAM KHẢO', ['PB','PE']),
]


def chart_image(pair,kind):
    """Mỗi biểu đồ là một ảnh riêng toàn chiều rộng, dùng chung số liệu với UI."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    font='Lato' if any(f.name=='Lato' for f in font_manager.fontManager.ttflist) else 'DejaVu Sans'
    m=pair['Metric'];scale=100 if m in PCT else 1 if m in MULT else 1e-9
    unit='%' if m in PCT else 'lần' if m in MULT else 'tỷ đồng'
    h,p=pair['History'],pair['Peers'];ph=pair.get('PeerHistory',pd.DataFrame())
    with plt.rc_context({'font.family':font,'font.size':9}):
        fig,ax=plt.subplots(figsize=(6.9,1.85 if kind=='history' else 1.65),dpi=190)
        if kind=='history':
            if len(h):
                vals=pd.to_numeric(h.Value,errors='coerce').where(h.Status.eq('OK'))*scale
                cuts=[0]+pair.get('HistoryBreaks',[])+[len(h)]
                for start,stop in zip(cuts,cuts[1:]):
                    ax.plot(range(start,stop),vals.iloc[start:stop],'o-',color='#2F6F12',lw=1.6,markersize=3,
                            label='Doanh nghiệp' if start==0 else '_nolegend_')
                ax.set_xticks(range(len(h)),h.DisplayPeriod,fontsize=9)
            if len(ph):
                cuts=[0]+pair.get('PeerHistoryBreaks',[])+[len(ph)]
                for start,stop in zip(cuts,cuts[1:]):
                    ax.plot(range(start,stop),ph.PeerMean.iloc[start:stop]*scale,'s--',color='#C07820',lw=1.4,markersize=3,
                            label='Bình quân peer' if start==0 else '_nolegend_')
            if not len(h) or not h.Value.notna().any():
                ax.text(.5,.5,'Chưa có chuỗi lịch sử doanh nghiệp',ha='center',va='center',transform=ax.transAxes)
            if len(h) or len(ph):ax.legend(loc='lower left',bbox_to_anchor=(0,1.02),ncol=2,fontsize=8,frameon=False,borderaxespad=0)
            title='Chuỗi năm gần nhất · Doanh nghiệp và bình quân peer'
            ax.margins(y=.22)
        else:
            z=included_rows(p)
            if len(z):
                ax.bar(z.Ticker,z.Value*scale,color=['#4F9E1E' if t==p.iloc[0].Ticker else '#ADB7A5' for t in z.Ticker],width=.6)
                ax.tick_params(axis='x',labelsize=9)
            else:
                ax.text(.5,.5,'Chưa có số cùng kỳ/cơ sở',ha='center',va='center',transform=ax.transAxes)
                ax.set_xticks([]);ax.set_yticks([])
            title=f"So sánh quý gần nhất · Doanh nghiệp và {pair['PeerCount']} peer · {pair['Period']}"
        ax.set_title(title,loc='left',fontsize=10,fontweight='bold',pad=24 if kind=='history' else 7)
        ax.set_ylabel(unit,fontsize=9);ax.tick_params(axis='y',labelsize=8)
        ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',alpha=.15)
        ax.ticklabel_format(axis='y',style='plain',useOffset=False)
        fig.tight_layout(pad=.6)
        stream=BytesIO();fig.savefig(stream,format='png');plt.close(fig);stream.seek(0)
    return stream


def pair_image(pair):
    """Tương thích các nơi gọi cũ: bản ghép cũng xếp dọc, không ghép ngang."""
    from PIL import Image
    images=[Image.open(chart_image(pair,k)).convert('RGB') for k in ('history','quarter')]
    canvas=Image.new('RGB',(max(i.width for i in images),sum(i.height for i in images)),'white')
    offset=0
    for image in images:canvas.paste(image,(0,offset));offset+=image.height
    stream=BytesIO();canvas.save(stream,format='PNG');stream.seek(0);return stream


def _para(doc,text,size=10.5,label=None,keep=False):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
    p.paragraph_format.widow_control=True;p.paragraph_format.keep_with_next=keep
    p.paragraph_format.keep_together=True
    if label:
        r=p.add_run(label+' ');r.bold=True;r.font.size=Pt(size)
    r=p.add_run(str(text));r.font.name='Lato';r.font.size=Pt(size)
    return p


def _table(doc,headers,rows,widths):
    from scripts.multisector_report import _set_cell_shading,_set_cell_margins
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    for i,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if i==0:tr.append(OxmlElement('w:tblHeader'))
        for j,(c,w) in enumerate(zip(row.cells,widths)):
            c.width=Mm(w);_set_cell_margins(c,top=30,start=65,bottom=30,end=65)
            _set_cell_shading(c,'E6F3D8' if i==0 else 'F6F8F3' if i%2==0 else 'FFFFFF')
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
                p.paragraph_format.keep_with_next=(i==0)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name='Lato';r.font.size=Pt(9);r.bold=i==0
    return t


def _cell_text(cell,text,bold=False,size=10):
    p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.05
    r=p.add_run(str(text));r.font.name='Lato';r.font.size=Pt(size);r.bold=bold
    return p


def _summary(doc,ticker,meta,rr,pack,report_type):
    from scripts.multisector_report import _section_band,_set_cell_shading,_set_cell_margins,company_display_name
    _section_band(doc,company_display_name(meta,ticker))
    cov=quantitative_coverage(ticker,pack);findings=executive_findings(ticker,pack)
    table=doc.add_table(rows=1,cols=2);table.autofit=False
    for col,w in zip(table.columns,[57,119]):col.width=Mm(w)
    left,right=table.rows[0].cells;left.width=Mm(57);right.width=Mm(119)
    _set_cell_shading(left,'E6F3D8')
    for c in [left,right]:_set_cell_margins(c,top=80,start=100,bottom=80,end=100)
    left.paragraphs[0].text='KẾT QUẢ XẾP HẠNG' if report_type=='rating' else 'THÔNG TIN PHÂN TÍCH'
    right.paragraphs[0].text='LUẬN ĐIỂM XẾP HẠNG' if report_type=='rating' else 'LUẬN ĐIỂM PHÂN TÍCH'
    for c in [left,right]:
        for r in c.paragraphs[0].runs:r.bold=True;r.font.size=Pt(11);r.font.name='Lato'
    _cell_text(left,rating_status_text(rr),True,11)
    for label,key in [('Anchor','Anchor'),('SCA','SCA'),('ICR','ICR')]:_cell_text(left,label+': '+rating_value_text(rr,key),size=10)
    _cell_text(left,'Triển vọng: '+str(rr.get('Outlook') or 'Chưa xác định'),size=10)
    _cell_text(left,'Kỳ tài chính: '+str(pack[0]['Period'] or 'Chưa có'),size=10)
    _cell_text(left,'Dữ liệu có số: '+str(len(cov['Present']))+'/'+str(cov['Total'])+' chỉ tiêu',size=10)
    for text in findings['Narrative']:_cell_text(right,text,size=10.5)
    if rr.get('RatingBlocker')=='SCOPE':
        _cell_text(right,rr['ApplicabilityReason'],size=10.5)
    else:
        _cell_text(right,'Kết luận tài chính dựa trên bằng chứng nêu trên. Bậc mô phỏng cần thêm hồ sơ định tính, dự báo vốn, thanh khoản và quyết định điều chỉnh theo phương pháp.',size=10)
    _cell_text(right,f'Nhóm tham chiếu ngành gồm {len(peer_universe(ticker))} peer: '+(', '.join(peer_universe(ticker)) or 'chưa có')+'.',size=10)
    _section_band(doc,'NHỮNG NHÂN TỐ CHÍNH DẪN ĐẾN KẾT QUẢ XẾP HẠNG' if report_type=='rating' else 'ĐIỂM MẠNH VÀ RỦI RO CẦN THEO DÕI')
    table=doc.add_table(rows=1,cols=2);table.autofit=False
    for col in table.columns:col.width=Mm(88)
    for idx,(head,items) in enumerate([('ĐIỂM MẠNH',findings['Strengths']),('ĐIỂM HẠN CHẾ / RỦI RO',findings['Risks'])]):
        c=table.cell(0,idx);c.width=Mm(88);_set_cell_margins(c,top=70,start=100,bottom=70,end=100)
        _set_cell_shading(c,'F5FAF0' if idx==0 else 'FAF7F2')
        c.paragraphs[0].text=head
        for r in c.paragraphs[0].runs:r.bold=True;r.font.size=Pt(10)
        for text in items:_cell_text(c,text,size=10)
    _para(doc,'Các điểm mạnh và hạn chế trên phản ánh số liệu quan sát, chưa phải điểm chấm tín nhiệm. '+cov['Note'],9)


def _context(doc,ticker,meta,pack,basis,years):
    from scripts.multisector_report import _section_band,_subhead,_researched_intel
    from scripts.company_profile import load_company_profile,profile_overview_text,profile_source_text
    _section_band(doc,'THÔNG TIN TỔNG QUAN TỔ CHỨC PHÁT HÀNH')
    profile=load_company_profile(ticker)
    if profile:
        _para(doc,profile_overview_text(ticker));_para(doc,profile_source_text(ticker),9)
    else:
        _para(doc,str(meta.get('DisplayName',ticker))+' ('+ticker+') thuộc nhóm so sánh '+SUBTYPE_LABELS.get(meta.get('InsuranceSubtype'),'chưa xác nhận')+'. Phạm vi số liệu: '+str(meta.get('InsuranceScope','chưa xác nhận'))+'.')
        _para(doc,'Nguồn phân loại trong hồ sơ: '+str(meta.get('InsuranceSource','chưa có'))+'.',9)
    if insurance_reporting_subtype(meta)=='GROUP':
        _para(doc,'BVH và PVI được xếp vào nhóm so sánh Phi nhân thọ theo cấu hình phân tích. Báo cáo vẫn sử dụng số liệu hợp nhất; cơ cấu nhân thọ, phi nhân thọ, tái bảo hiểm và đầu tư có thể khác nhau. Đối chiếu mang tính tham khảo, không thay số liệu riêng của pháp nhân phi nhân thọ để chấm XHTN.')
    for title,kind,fallback in [
        ('RỦI RO VĨ MÔ','MACRO','Lãi suất tác động đến thu nhập tái đầu tư và giá trị danh mục; lạm phát chi phí có thể làm tăng mức bồi thường. Khi thẩm định cần đối chiếu thời hạn tài sản với nghĩa vụ bảo hiểm, độ nhạy lợi nhuận và vốn trong kịch bản bất lợi.'),
        ('RỦI RO NGÀNH','INDUSTRY','Cạnh tranh phí, biến động tổn thất và điều kiện tái bảo hiểm ảnh hưởng kết quả khai thác. Tăng doanh thu cần được đọc cùng bồi thường, chi phí, dự phòng và chất lượng đối tác tái bảo hiểm; quy mô lớn chưa đủ kết luận rủi ro thấp.')]:
        _section_band(doc,title)
        research=_researched_intel(ticker,meta,kind,limit=2)
        if research:
            for name,narrative,source,url,asof in research:
                if name:_subhead(doc,name)
                _para(doc,narrative);_para(doc,'Nguồn: '+source+'; kỳ nguồn: '+asof+(' — '+url if url else ''),9)
        else:
            _para(doc,fallback)
            _para(doc,'Phần này nêu cơ chế rủi ro để phân tích; chưa có hồ sơ nghiên cứu theo kỳ để gán điểm vĩ mô/ngành.',9)
    _section_band(doc,'CƠ SỞ PHÂN TÍCH VÀ SO SÁNH')
    _para(doc,f'Lịch sử: {BASIS_LABELS[basis]}, tối đa {years} năm gần nhất có trong phạm vi yêu cầu. So sánh ngang dùng {pack[0]["Period"] or "quý chưa xác định"}; không chuyển sang quý khác khi peer thiếu số. Thu nhập/chi phí là quý riêng lẻ, số dư là cuối quý; ROE, ROA, P/B và P/E giữ cơ sở tỷ số nguồn.',10)
    _para(doc,POLICY_NOTE+' So sánh quý công bố bình quân và trung vị khi đủ ít nhất 5 peer cùng kỳ, cùng cơ sở.',10)
    from scripts.insurance_history import HISTORY_PEER_NOTE
    _para(doc,'Đường bình quân lịch sử: '+HISTORY_PEER_NOTE,10)
    for subtype,label in SUBTYPE_LABELS.items():
        members=[t for t in [ticker]+peer_universe(ticker) if get_company(t).get('InsuranceSubtype')==subtype]
        if members:_para(doc,label+': '+', '.join(members)+'.',10)


def _metric(doc,ticker,pair):
    from scripts.multisector_report import _subhead
    m=pair['Metric'];h,p=pair['History'],pair['Peers']
    _subhead(doc,pair['Label'])
    chart=doc.add_paragraph();chart.paragraph_format.space_after=Pt(2);chart.paragraph_format.keep_with_next=True
    chart.add_run().add_picture(chart_image(pair,'history'),width=Mm(176))
    if len(h):
        ph=pair['PeerHistory']
        rows=[[ticker]+[fmt(m,r.Value,r.Status,True) if r.Status=='OK' else '—' for r in h.itertuples()],
              ['TB peer']+[fmt(m,r.PeerMean,'OK',True) if pd.notna(r.PeerMean) else '—' for r in ph.itertuples()],
              ['Số peer (N)']+[str(r.PeerCount) for r in ph.itertuples()]]
        _table(doc,['Lịch sử']+h.DisplayPeriod.tolist(),rows,[26]+[150/len(h)]*len(h))
    # Chuỗi có từng năm và thay đổi năm liền kề; giữ nguyên cơ sở UI.
    history=pair['NarrativeHistory']
    _para(doc,history,10,label='Diễn biến nhiều năm:')
    for warning in pair.get('HistoryWarnings',[]):_para(doc,warning,9,label='Lưu ý chất lượng dữ liệu:')
    chart=doc.add_paragraph();chart.paragraph_format.space_after=Pt(2);chart.paragraph_format.keep_with_next=True
    chart.add_run().add_picture(chart_image(pair,'quarter'),width=Mm(176))
    _para(doc,pair['NarrativePeer'],10,label='So sánh quý gần nhất:')
    _para(doc,pair['AnalysisText'],10,label='Hàm ý phân tích:',keep=True)
    _para(doc,pair['ConclusionText'],10,label='Kết luận:')


def _qualitative(doc,rr,pack,subtype):
    from scripts.multisector_report import _section_band
    _section_band(doc,'KẾT LUẬN PHÂN TÍCH')
    for text in overall_conclusion(pack,subtype):_para(doc,text)
    _section_band(doc,'QUẢN TRỊ VÀ QUẢN LÝ')
    _para(doc,'Đánh giá quản trị cần xem cơ cấu kiểm soát, khẩu vị rủi ro, chính sách tài chính và chất lượng thực thi. Số liệu lợi nhuận hoặc quy mô không thay thế bằng chứng về các yếu tố này.')
    reviewed=[r for r in rr.get('SubfactorEvidence',[]) if r.get('Đủ bằng chứng') and r.get('Nhóm','').startswith('V.')]
    if reviewed:
        for r in reviewed:_para(doc,r['Chỉ tiêu']+': '+r['Đánh giá']+'. '+r['Lập luận']+' Nguồn: '+r['Nguồn'],10)
    elif rr.get('RatingBlocker')!='SCOPE':
        _para(doc,'Hồ sơ hiện chưa có đánh giá định tính đủ nguồn và lập luận. Các nhóm hồ sơ cần bổ sung được tổng hợp cuối báo cáo.',10)
    _section_band(doc,'YẾU TỐ BÊN NGOÀI')
    o=rr.get('AssessmentInputs',{})
    if o.get('SupportSource'):_para(doc,str(o['SupportSource']))
    else:_para(doc,'Chưa có hồ sơ xác nhận năng lực, động cơ và cơ chế hỗ trợ bên ngoài. Không mặc định quan hệ cổ đông hoặc thương hiệu tạo ra hỗ trợ; chưa xác định số bậc điều chỉnh.')
    _section_band(doc,'ĐỘ NHẠY XẾP HẠNG')
    sensitivity=sensitivity_narrative(pack,subtype)
    _para(doc,sensitivity['Positive'],label='Điều kiện cải thiện:')
    _para(doc,sensitivity['Negative'],label='Điều kiện gây sức ép:')
    _para(doc,sensitivity['Limit'],9.5)
    _section_band(doc,'TỔNG KẾT CẤU PHẦN XẾP HẠNG')
    _para(doc,'Các kết luận về kinh doanh, sinh lời, vốn và thanh khoản ở trên là cơ sở lập luận của báo cáo. Trạng thái dưới đây phản ánh bước thẩm định và phạm vi áp dụng phương pháp, không xóa bỏ các kết luận tài chính đã có bằng chứng.',10)
    _table(doc,['Cấu phần','Kết quả','Cơ sở'],[[r['Cấu phần'],r['Kết quả'],r.get('Luận cứ','')] for r in insurance_scorecard(rr)],[48,43,85])
    if o.get('DecisionReason'):_para(doc,str(o['DecisionReason']))


def _appendix(doc,pack,rr):
    from scripts.multisector_report import _section_band,_subhead
    _section_band(doc,'PHỤ LỤC — DỮ LIỆU VÀ HỒ SƠ CẦN BỔ SUNG')
    missing=[];na=[]
    for p in pack:
        q=p['Peers'];r=q.iloc[0] if len(q) else None
        status=r.Status if r is not None else 'MISSING'
        if status=='OK':continue
        reason=str(r.Reason) if r is not None else 'Chưa có kỳ dữ liệu tài chính.'
        reason=data_gap_explanation(p['Metric'],status,reason)
        (na if status=='NOT_APPLICABLE' else missing).append([p['Label'],STATUS.get(status,status),reason])
    if missing:
        _subhead(doc,'Chỉ tiêu cần bổ sung nguồn đúng kỳ')
        _table(doc,['Chỉ tiêu','Trạng thái','Lý do / hồ sơ cần có'],missing,[57,28,91])
    if na:
        _subhead(doc,'Chỉ tiêu không áp dụng ở phạm vi đang phân tích')
        _table(doc,['Chỉ tiêu','Trạng thái','Lý do'],na,[57,28,91])
    pending=pending_assessment_groups(rr)
    if pending:
        _subhead(doc,'Nhóm hồ sơ thẩm định còn thiếu')
        _para(doc,f"Đầu vào thẩm định đã có: {rr.get('AssessmentInputCount',0)}/{rr.get('AssessmentInputTotal',0)}. Phiếu chi tiết trên ứng dụng lưu từng đánh giá, lập luận và nguồn; bảng dưới chỉ tổng hợp theo nhóm.",10)
        _table(doc,['Nhóm hồ sơ','Có bằng chứng','Nội dung cần bổ sung'],[[r['Nhóm hồ sơ'],r['Có bằng chứng'],r['Cần bổ sung']] for r in pending],[48,27,101])
    _section_band(doc,'NGUỒN, CÔNG THỨC VÀ GIỚI HẠN PHƯƠNG PHÁP')
    _para(doc,'BCTC và tỷ số nguồn: bộ dữ liệu đã nạp trong project, đối chiếu theo kỳ của từng mã. Chưa xác nhận kiểm toán hoặc tình trạng soát xét của từng kỳ chỉ từ tệp số liệu. Dấu “—” chỉ thiếu dữ liệu hoặc không cùng cơ sở, không phải số 0.',9.5)
    _para(doc,'Tham chiếu: Phương pháp XHTN Công ty Bảo hiểm Phi Nhân thọ, Saigon Ratings, cập nhật 20.01.2025, tài liệu do người dùng cung cấp. Quy tắc định tính và dự báo cần bằng chứng thẩm định; không gán điểm mặc định hoặc dùng nhận xét tự động thay quyết định Hội đồng.',9.5)
    for p in pack:
        if not has_evidence(p):continue
        q=p['Peers'];src=str(q.iloc[0].Source) if len(q) else ''
        _para(doc,p['Label']+': '+p['MethodNote']+(' Nguồn/công thức: '+src if src else ''),9)
    _section_band(doc,'THÀNH PHẦN PEER TRONG CHUỖI NĂM')
    _para(doc,'Bình quân trong biểu đồ và bảng lịch sử dùng các mã dưới đây tại đúng từng kỳ. Thành phần mẫu hiện được chọn theo danh mục phân ngành của project.',9)
    for p in pack:
        if not has_evidence(p):continue
        ph=p['PeerHistory'];groups={}
        for r in ph.itertuples():groups.setdefault(tuple(r.PeerTickers),[]).append(r.DisplayPeriod)
        details=[]
        for members,periods in groups.items():
            details.append(', '.join(periods)+': '+(', '.join(members) if members else 'chưa có peer hợp lệ'))
        _para(doc,p['Label']+' — '+'; '.join(details)+'.',9)
    _section_band(doc,'CHI TIẾT MẪU SO SÁNH QUÝ GẦN NHẤT')
    for p in pack:
        if not has_evidence(p):continue
        _subhead(doc,p['Label'])
        q=p['Peers'];m=p['Metric']
        values=[f"{r.Ticker}: {fmt(m,r.Value,r.Status)}" for r in q.itertuples() if r.Included]
        _para(doc,'; '.join(values)+'.',9)
        if p['PeerExclusions']:_para(doc,p['PeerExclusions'],9)


def generate_insurance_docx(ticker,report_type='analysis',rating_result=None,options=None):
    from scripts.multisector_report import _style_doc,_cover_sample,_section_band
    ticker=str(ticker).upper();options=options or {}
    basis=options.get('history_basis','FY');years=int(options.get('history_years',5))
    meta=get_company(ticker);pack=analysis_pack(ticker,basis,years);rr=rating_result or insurance_rating(ticker)
    subtype=insurance_reporting_subtype(meta)
    findings=section_findings(pack,subtype)
    doc=Document();_style_doc(doc,report_type)
    doc.styles['Normal'].font.size=Pt(10.5)
    _cover_sample(doc,ticker,meta,report_type)
    _summary(doc,ticker,meta,rr,pack,report_type)
    doc.add_page_break();_context(doc,ticker,meta,pack,basis,years)
    for row in rr.get('SubfactorEvidence',[]):
        if row.get('Đủ bằng chứng') and row.get('Nhóm','').startswith('I.'):
            _para(doc,row['Lập luận']+' Nguồn: '+row['Nguồn'],10,label=row['Chỉ tiêu']+' — '+row['Đánh giá']+':')
    pairs={p['Metric']:p for p in pack}
    for heading,metrics in SECTIONS:
        available=[pairs[m] for m in metrics if has_evidence(pairs[m])]
        if heading == 'HỒ SƠ KINH DOANH':
            doc.add_page_break()
        _section_band(doc,heading)
        from scripts.multisector_report import _subhead
        for title,text in findings[heading]:
            _subhead(doc,title);_para(doc,text)
        evidence_groups={'HỒ SƠ KINH DOANH':['II.'],'VỐN VÀ DỰ PHÒNG NGHIỆP VỤ':['III.'],
                         'THANH KHOẢN VÀ KHẢ NĂNG THANH TOÁN':['IV.']}.get(heading,[])
        for row in rr.get('SubfactorEvidence',[]):
            if row.get('Đủ bằng chứng') and any(row.get('Nhóm','').startswith(x) for x in evidence_groups):
                _para(doc,row['Lập luận']+' Nguồn: '+row['Nguồn'],10,label=row['Chỉ tiêu']+' — '+row['Đánh giá']+':')
        for pair in available:
            _metric(doc,ticker,pair)
        if report_type=='analysis' and heading=='ĐỊNH GIÁ THAM KHẢO':
            from scripts.multisector_valuation import valuation
            from scripts.fair_value_range import fair_value_range
            from scripts.universal_data import get_snapshot
            v=valuation(ticker,get_snapshot(ticker));fv=fair_value_range(ticker)
            _para(doc,str(v.get('PrimaryMethod',''))+'. '+str(v.get('Reason','')),10)
            _table(doc,['Kịch bản','Giá trị đồng/cổ phiếu'],[[k,fmt('FairValue',fv.get(k),v.get('Status'))] for k in ['Bear','Base','Bull','StrategicMA']],[88,88])
    _qualitative(doc,rr,pack,subtype)
    _appendix(doc,pack,rr)
    sec=doc.sections[0];footer=sec.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run(ticker+' | Lê Hoàng Quân | V'+ENGINE_VERSION+' | '+date.today().strftime('%d/%m/%Y')+' | Trang ');r.font.name='Lato';r.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for scope in [doc.element,doc.styles.element]:
        for el in list(scope.iter(qn('w:pBdr'))):el.getparent().remove(el)
        for el in scope.iter(qn('w:color')):el.attrib.pop(qn('w:themeColor'),None)
    # Ngắt trang gắn vào đoạn kế tiếp, tránh trang trắng khi đoạn trước vừa đầy trang.
    from docx.text.paragraph import Paragraph
    for node in list(doc.element.body):
        if node.tag!=qn('w:p'):continue
        paragraph=Paragraph(node,doc)
        breaks=node.findall('.//'+qn('w:br'))
        nxt=node.getnext()
        if not paragraph.text.strip() and any(b.get(qn('w:type'))=='page' for b in breaks) and nxt is not None and nxt.tag==qn('w:p'):
            following=Paragraph(nxt,doc)
            following.paragraph_format.page_break_before=True
            following.paragraph_format.keep_with_next=True
            node.getparent().remove(node)
    out=BytesIO();doc.save(out);return out.getvalue()

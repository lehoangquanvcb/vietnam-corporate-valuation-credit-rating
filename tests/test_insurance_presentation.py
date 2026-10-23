"""Hồi quy cho lỗi hiển thị BVH và bố cục báo cáo do người dùng phản ánh."""
import unittest
from io import BytesIO
from docx import Document
from scripts.insurance_rating import insurance_rating,insurance_scorecard,rating_value_text,rating_status_text
from scripts.insurance_presentation import quantitative_coverage,pending_assessment_groups
from scripts.insurance_analysis import metric_pair,analysis_pack
from scripts.rating_committee_engine import committee_pack
from scripts.data_quality_engine import assess_data_quality
from tests.test_insurance_workflow import assessed_inputs


def document_text(doc):
    return '\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])


class Presentation(unittest.TestCase):
    def test_out_of_scope_is_not_pending_assessment(self):
        for ticker in ['BVH','PVI','PRE','VNR']:
            r=insurance_rating(ticker)
            self.assertEqual(r['RatingBlocker'],'SCOPE')
            self.assertEqual(r['ICR'],'N/A')
            self.assertEqual(rating_value_text(r),'Không áp dụng')
            self.assertEqual(r['AssessmentInputTotal'],0)
            self.assertEqual(pending_assessment_groups(r),[])
            self.assertLessEqual(len(insurance_scorecard(r)),2)
            self.assertNotIn('Chưa thẩm định',str(committee_pack(ticker)['Waterfall']))

    def test_unreviewed_nonlife_keeps_gates_without_placeholder_wall(self):
        r=insurance_rating('ABI');rows=insurance_scorecard(r)
        self.assertEqual(r['ICR'],'N/A');self.assertFalse(r['AssessmentComplete'])
        self.assertEqual(rating_value_text(r),'Chưa xác định')
        self.assertEqual(len(rows),1)
        self.assertGreater(len(pending_assessment_groups(r)),0)
        self.assertEqual(r['AssessmentInputTotal'],65)

    def test_matrix_na_distinct_from_scope(self):
        r=insurance_rating('ABI',assessed_inputs(CapitalAdequacy=8,FinancialAdjustment=1))
        self.assertEqual(r['RatingBlocker'],'MATRIX')
        self.assertEqual(rating_status_text(r),'Ma trận không quy định bậc')
        self.assertNotEqual(rating_value_text(r),'Không áp dụng')
        self.assertGreater(len(insurance_scorecard(r)),1)

    def test_partial_input_does_not_discard_calculated_sca(self):
        r=insurance_rating('ABI',assessed_inputs(ExternalSupportNotches=None))
        self.assertEqual(rating_value_text(r,'SCA'),'AA+')
        self.assertEqual(rating_value_text(r,'ICR'),'Chưa xác định')
        rows=insurance_scorecard(r)
        self.assertTrue(any(row['Cấu phần']=='SCA — xếp hạng độc lập' and row['Kết quả']=='vnAA+' for row in rows))

    def test_coverage_accounts_for_all_34_metrics(self):
        d=assess_data_quality('BVH')
        self.assertEqual((len(d['Present']),len(d['Missing']),len(d['NotApplicable'])),(18,11,5))
        self.assertEqual(d['Total'],34);self.assertEqual(d['ApplicableTotal'],29)
        self.assertAlmostEqual(d['Coverage'],18/29)
        self.assertEqual(d,quantitative_coverage('BVH',analysis_pack('BVH')))
        a=assess_data_quality('ABI')
        self.assertEqual((len(a['Present']),len(a['Missing']),len(a['NotApplicable'])),(22,12,0))

    def test_peer_narrative_names_actual_same_quarter_peer(self):
        p=metric_pair('BVH','ROE')
        self.assertIn('PVI',set(p['Peers'].Ticker));self.assertIn('2026-Q2',p['PeerText'])
        self.assertEqual(p['PeerCount'],11);self.assertIsNotNone(p['PeerMean'])
        self.assertIn('Biến động từng năm',p['HistoryText'])
        self.assertEqual(set(p['Peers'].Period),{'2026-Q2'})
        self.assertIn('không diễn giải',metric_pair('BVH','UnderwritingProfit')['AnalysisText'].lower())

    def test_reserve_source_break_is_not_interpreted_as_growth(self):
        p=metric_pair('BVH','InsuranceReserves')
        self.assertEqual(p['HistoryBreaks'],[1]);self.assertTrue(p['HistoryWarnings'])
        self.assertIn('Chưa tính thay đổi toàn chuỗi',p['HistoryText'])
        self.assertNotIn('70669',p['HistoryText'])
        self.assertEqual(p['History'].iloc[0].Value,285420048108)


class WordReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from scripts.evidence_report import generate_docx
        cls.doc=Document(BytesIO(generate_docx('BVH','rating',mna={'history_basis':'FY','history_years':5})))
        cls.text=document_text(cls.doc)

    def test_report_matches_reference_structure(self):
        for label in ['KẾT QUẢ XẾP HẠNG','LUẬN ĐIỂM XẾP HẠNG','ĐIỂM MẠNH','ĐIỂM HẠN CHẾ / RỦI RO',
                      'THÔNG TIN TỔNG QUAN TỔ CHỨC PHÁT HÀNH','RỦI RO VĨ MÔ','RỦI RO NGÀNH','HỒ SƠ KINH DOANH',
                      'VỐN VÀ DỰ PHÒNG NGHIỆP VỤ','THANH KHOẢN VÀ KHẢ NĂNG THANH TOÁN','QUẢN TRỊ VÀ QUẢN LÝ',
                      'YẾU TỐ BÊN NGOÀI','ĐỘ NHẠY XẾP HẠNG','TỔNG KẾT CẤU PHẦN XẾP HẠNG']:
            self.assertIn(label,self.text)
        self.assertEqual(self.doc.styles['Normal'].font.name,'Lato')
        self.assertAlmostEqual(self.doc.sections[0].page_width.mm,210,places=1)
        self.assertIn('4F9E1E',self.doc.element.xml)

    def test_only_evidenced_metrics_receive_chart_pairs(self):
        self.assertEqual(len(self.doc.inline_shapes),18)
        self.assertEqual(self.text.count('Diễn biến nhiều năm:'),18)
        self.assertEqual(self.text.count('So sánh quý gần nhất:'),18)
        self.assertEqual(self.text.count('Hàm ý phân tích:'),18)
        self.assertIn('PHỤ LỤC — DỮ LIỆU VÀ HỒ SƠ CẦN BỔ SUNG',self.text)
        self.assertNotIn('Chưa thẩm định',self.text)
        self.assertNotIn('0/65',self.text)
        self.assertIn('Ngoài phạm vi phương pháp',self.text)
        for pair in analysis_pack('BVH'):
            self.assertIn(pair['Label'],self.text)


if __name__=='__main__':unittest.main()
